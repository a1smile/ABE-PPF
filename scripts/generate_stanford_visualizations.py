import argparse
import csv
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.append(str(REPO_ROOT))

from scripts.build_paper_outputs import (  # noqa: E402
    METHOD_GROUP,
    STANFORD_BATCHES,
    load_batch_records,
    load_points,
    sample_points,
    transform_points,
)


OURS_METHOD = "Ours (Full)"
DPI = 320
SUCCESS_THRESHOLD = 0.02
VIEWS = [
    ("front", 0, -90),
    ("back", 0, 90),
    ("left", 0, 180),
    ("right", 0, 0),
    ("top", 90, -90),
    ("bottom", -90, -90),
    ("iso_left", 24, 45),
    ("iso_right", 24, -45),
]
FIELDNAMES = [
    "group",
    "method",
    "model_label",
    "model_slug",
    "obj_id",
    "case_id",
    "idx",
    "scene_name",
    "scene_variant",
    "primary_success",
    "ADD",
    "ADD_S",
    "rotation_error_deg",
    "translation_error",
    "registration_time",
    "paired_ours_ADD",
    "paired_ours_rotation_error_deg",
    "paired_ours_translation_error",
    "selection_reason",
    "status",
    "note",
    "view_name",
    "output_path",
]


def slugify(text: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")
    return slug or "unknown"


def safe_float(value: Any, default: float = float("inf")) -> float:
    try:
        out = float(value)
    except (TypeError, ValueError):
        return default
    if not np.isfinite(out):
        return default
    return out


def load_model_mapping(path: Path) -> tuple[dict[str, str], dict[int, str]]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    by_filename: dict[str, str] = {}
    by_obj_id: dict[int, str] = {}
    for model_filename, model_label in ws.iter_rows(min_row=2, values_only=True):
        if not model_filename or not model_label:
            continue
        filename = str(model_filename)
        label = str(model_label)
        by_filename[filename] = label
    return by_filename, by_obj_id


def build_obj_id_to_label(
    render_lookup: dict[tuple[str, str], dict[str, Any]],
    filename_to_label: dict[str, str],
) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for (_, _case_id), panel in render_lookup.items():
        obj_id = panel.get("obj_id")
        model_path = panel.get("model_path")
        if obj_id is None or model_path is None:
            continue
        filename = Path(str(model_path)).name
        label = filename_to_label.get(filename)
        if label:
            mapping[int(obj_id)] = label
    return mapping


def build_records() -> tuple[list[dict[str, Any]], dict[tuple[str, str], dict[str, Any]]]:
    rows: list[dict[str, Any]] = []
    render_lookup: dict[tuple[str, str], dict[str, Any]] = {}
    for method, batch_path in STANFORD_BATCHES.items():
        if not batch_path.exists():
            continue
        method_rows, render_rows = load_batch_records("Stanford", method, batch_path)
        rows.extend(method_rows)
        for item in render_rows:
            render_lookup[(method, item["case_id"])] = item
    return rows, render_lookup


def infer_model_label(record: dict[str, Any], filename_to_label: dict[str, str], obj_id_to_label: dict[int, str]) -> str:
    obj_id = record.get("obj_id")
    if obj_id is not None and int(obj_id) in obj_id_to_label:
        return obj_id_to_label[int(obj_id)]
    row = record.get("row", {}) or {}
    model_name = row.get("model_name")
    if model_name and model_name in filename_to_label:
        return filename_to_label[model_name]
    model_path = record.get("model_path")
    if model_path:
        name = Path(str(model_path)).name
        if name in filename_to_label:
            return filename_to_label[name]
    return model_name or f"obj_{record.get('obj_id')}"


def render_case(panel: dict[str, Any], out_path: Path, elev: float, azim: float) -> None:
    scene_pts = load_points(panel["scene_path"])
    model_pts = load_points(panel["model_path"])
    T_pred = np.asarray(panel["T_pred"], dtype=np.float64)
    if scene_pts.size == 0 or model_pts.size == 0 or T_pred.size == 0:
        raise ValueError("missing point cloud or transform data")

    scene_vis = sample_points(scene_pts, 4200, seed=0)
    model_vis = sample_points(transform_points(model_pts, T_pred), 2000, seed=1)
    all_pts = np.concatenate([scene_vis, model_vis], axis=0)

    fig = plt.figure(figsize=(5.4, 5.4), facecolor="black")
    ax = fig.add_subplot(111, projection="3d")
    ax.set_facecolor("black")
    ax.scatter(scene_vis[:, 0], scene_vis[:, 1], scene_vis[:, 2], s=0.72, c="#D9DEE5", alpha=0.46, linewidths=0)
    ax.scatter(model_vis[:, 0], model_vis[:, 1], model_vis[:, 2], s=1.35, c="#FF6A55", alpha=0.98, linewidths=0)
    ax.view_init(elev=elev, azim=azim)
    mins = all_pts.min(axis=0)
    maxs = all_pts.max(axis=0)
    center = (mins + maxs) / 2.0
    radius = max(float(np.max(maxs - mins)) / 2.0, 1e-6) * 0.70
    ax.set_xlim(center[0] - radius, center[0] + radius)
    ax.set_ylim(center[1] - radius, center[1] + radius)
    ax.set_zlim(center[2] - radius, center[2] + radius)
    ax.set_axis_off()
    fig.subplots_adjust(left=0.0, right=1.0, bottom=0.0, top=1.0)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=DPI, facecolor="black", edgecolor="none")
    plt.close(fig)


def ours_best_key(rec: dict[str, Any]) -> tuple[Any, ...]:
    return (
        0 if rec.get("primary_success") else 1,
        safe_float(rec.get("ADD")),
        safe_float(rec.get("rotation_error_deg")),
        safe_float(rec.get("translation_error")),
        safe_float(rec.get("registration_time")),
        safe_float(rec.get("idx"), default=1e18),
    )


def external_bad_key(rec: dict[str, Any], ours_rec: dict[str, Any]) -> tuple[Any, ...]:
    return (
        safe_float(rec.get("ADD")) - safe_float(ours_rec.get("ADD"), default=0.0),
        safe_float(rec.get("rotation_error_deg")),
        safe_float(rec.get("translation_error")),
        safe_float(rec.get("registration_time")),
    )


def external_worst_success_key(rec: dict[str, Any], ours_rec: dict[str, Any]) -> tuple[Any, ...]:
    return (
        safe_float(rec.get("ADD")) - safe_float(ours_rec.get("ADD"), default=0.0),
        safe_float(rec.get("rotation_error_deg")) - safe_float(ours_rec.get("rotation_error_deg"), default=0.0),
        safe_float(rec.get("translation_error")) - safe_float(ours_rec.get("translation_error"), default=0.0),
        safe_float(rec.get("registration_time")) - safe_float(ours_rec.get("registration_time"), default=0.0),
    )


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--mapping_xlsx", required=True, type=Path)
    ap.add_argument("--output_root", default=REPO_ROOT / "visualizations" / "stanford", type=Path)
    args = ap.parse_args()

    filename_to_label, _ = load_model_mapping(args.mapping_xlsx)
    records, render_lookup = build_records()
    obj_id_to_label = build_obj_id_to_label(render_lookup, filename_to_label)

    ours_records = [r for r in records if r["method"] == OURS_METHOD]
    by_model_ours: dict[str, list[dict[str, Any]]] = defaultdict(list)
    ours_by_case: dict[str, dict[str, Any]] = {}
    for rec in ours_records:
        label = infer_model_label(rec, filename_to_label, obj_id_to_label)
        rec["model_label"] = label
        by_model_ours[label].append(rec)
        ours_by_case[rec["case_id"]] = rec

    rows_out: list[dict[str, Any]] = []
    ours_best_by_model: dict[str, dict[str, Any]] = {}

    for model_label, group in sorted(by_model_ours.items()):
        chosen = sorted(group, key=ours_best_key)[0]
        ours_best_by_model[model_label] = chosen
        slug = slugify(model_label)
        panel = render_lookup.get((OURS_METHOD, chosen["case_id"]))
        for view_name, elev, azim in VIEWS:
            out_path = args.output_root / "ours_best" / f"{slug}_ours_best_{view_name}.png"
            status = "saved"
            note = ""
            if panel is None:
                status = "missing"
                note = "No render panel found for selected case."
            else:
                try:
                    render_case(panel, out_path, elev=elev, azim=azim)
                except Exception as exc:
                    status = "missing"
                    note = f"Render failed: {exc}"
            rows_out.append(
                {
                    "group": "ours_best",
                    "method": OURS_METHOD,
                    "model_label": model_label,
                    "model_slug": slug,
                    "obj_id": chosen.get("obj_id"),
                    "case_id": chosen.get("case_id"),
                    "idx": chosen.get("idx"),
                    "scene_name": chosen.get("scene_name"),
                    "scene_variant": chosen.get("scene_variant"),
                    "primary_success": chosen.get("primary_success"),
                    "ADD": chosen.get("ADD"),
                    "ADD_S": chosen.get("ADD_S"),
                    "rotation_error_deg": chosen.get("rotation_error_deg"),
                    "translation_error": chosen.get("translation_error"),
                    "registration_time": chosen.get("registration_time"),
                    "paired_ours_ADD": "",
                    "paired_ours_rotation_error_deg": "",
                    "paired_ours_translation_error": "",
                    "selection_reason": "primary_success first, then minimum ADD, then RE/TE/time",
                    "status": status,
                    "note": note,
                    "view_name": view_name,
                    "output_path": str(out_path) if status == "saved" else "",
                }
            )

    external_methods = [m for m, group in METHOD_GROUP.items() if group == "external"]
    for method in external_methods:
        ext_records = [r for r in records if r["method"] == method]
        by_model_ext: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for rec in ext_records:
            label = infer_model_label(rec, filename_to_label, obj_id_to_label)
            rec["model_label"] = label
            by_model_ext[label].append(rec)

        for model_label in sorted(by_model_ours):
            slug = slugify(model_label)
            failure_candidates = []
            fallback_success_candidates = []
            for rec in by_model_ext.get(model_label, []):
                ours_rec = ours_by_case.get(rec["case_id"])
                if ours_rec is None:
                    continue
                if bool(ours_rec.get("primary_success")) and not bool(rec.get("primary_success")):
                    failure_candidates.append((rec, ours_rec))
                elif bool(rec.get("primary_success")):
                    fallback_success_candidates.append((rec, ours_rec))

            selection_reason = ""
            if failure_candidates:
                chosen_rec, paired_ours = max(failure_candidates, key=lambda item: external_bad_key(item[0], item[1]))
                selection_reason = "Ours succeeds while external fails; maximize ADD gap against Ours"
            elif fallback_success_candidates:
                chosen_rec, paired_ours = max(
                    fallback_success_candidates,
                    key=lambda item: external_worst_success_key(item[0], item[1]),
                )
                selection_reason = "No failure case found; use the worst successful external case by degradation against Ours"
            else:
                rows_out.append(
                    {
                        "group": "baseline_bad_cases",
                        "method": method,
                        "model_label": model_label,
                        "model_slug": slug,
                        "obj_id": "",
                        "case_id": "",
                        "idx": "",
                        "scene_name": "",
                        "scene_variant": "",
                        "primary_success": "",
                        "ADD": "",
                        "ADD_S": "",
                        "rotation_error_deg": "",
                        "translation_error": "",
                        "registration_time": "",
                        "paired_ours_ADD": "",
                        "paired_ours_rotation_error_deg": "",
                        "paired_ours_translation_error": "",
                        "selection_reason": "prefer failure cases; otherwise worst successful case",
                        "status": "skipped",
                        "note": "No failure case or usable successful case found for this method/model.",
                        "view_name": "",
                        "output_path": "",
                    }
                )
                continue

            panel = render_lookup.get((method, chosen_rec["case_id"]))
            for view_name, elev, azim in VIEWS:
                out_path = (
                    args.output_root
                    / "baseline_bad_cases"
                    / slugify(method)
                    / f"{slug}_{slugify(method)}_bad_{view_name}.png"
                )
                status = "saved"
                note = ""
                if panel is None:
                    status = "missing"
                    note = "No render panel found for selected case."
                else:
                    try:
                        render_case(panel, out_path, elev=elev, azim=azim)
                    except Exception as exc:
                        status = "missing"
                        note = f"Render failed: {exc}"
                rows_out.append(
                    {
                        "group": "baseline_bad_cases",
                        "method": method,
                        "model_label": model_label,
                        "model_slug": slug,
                        "obj_id": chosen_rec.get("obj_id"),
                        "case_id": chosen_rec.get("case_id"),
                        "idx": chosen_rec.get("idx"),
                        "scene_name": chosen_rec.get("scene_name"),
                        "scene_variant": chosen_rec.get("scene_variant"),
                        "primary_success": chosen_rec.get("primary_success"),
                        "ADD": chosen_rec.get("ADD"),
                        "ADD_S": chosen_rec.get("ADD_S"),
                        "rotation_error_deg": chosen_rec.get("rotation_error_deg"),
                        "translation_error": chosen_rec.get("translation_error"),
                        "registration_time": chosen_rec.get("registration_time"),
                        "paired_ours_ADD": paired_ours.get("ADD"),
                        "paired_ours_rotation_error_deg": paired_ours.get("rotation_error_deg"),
                        "paired_ours_translation_error": paired_ours.get("translation_error"),
                        "selection_reason": selection_reason,
                        "status": status,
                        "note": note,
                        "view_name": view_name,
                        "output_path": str(out_path) if status == "saved" else "",
                    }
                )

    summary_path = args.output_root / "visualization_summary.csv"
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    with summary_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows_out)

    saved = sum(1 for r in rows_out if r["status"] == "saved")
    skipped = sum(1 for r in rows_out if r["status"] == "skipped")
    missing = sum(1 for r in rows_out if r["status"] == "missing")
    print(f"summary {summary_path}")
    print(f"saved {saved}")
    print(f"skipped {skipped}")
    print(f"missing {missing}")


if __name__ == "__main__":
    main()
