import argparse
import json
from collections import OrderedDict, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from nature_table_utils import export_table_bundle


COLUMNS = [
    {"header": "Object", "align": "left"},
    {"header": "N", "align": "right"},
    {"header": "Baseline SR", "align": "right"},
    {"header": "Ours SR", "align": "right"},
    {"header": "Delta SR", "align": "right"},
    {"header": "Ours ADD (mm)", "key": "Ours ADD", "align": "right"},
    {"header": "Ours ADD-S (mm)", "key": "Ours ADD-S", "align": "right"},
    {"header": "Ours RE (deg)", "key": "Ours RE", "align": "right"},
    {"header": "Ours TE (mm)", "key": "Ours TE", "align": "right"},
    {"header": "Ours Time (s)", "key": "Ours Time", "align": "right"},
]


def load_mapping(path: Path) -> OrderedDict[int, str]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    mapping: OrderedDict[int, str] = OrderedDict()
    for model_name, object_name in ws.iter_rows(min_row=2, values_only=True):
        if not model_name or not object_name:
            continue
        stem = Path(str(model_name)).stem
        obj_id = int(stem.split("_")[-1])
        mapping[obj_id] = str(object_name)
    return mapping


def load_results(path: Path) -> list[dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return payload.get("results", []) or []


def mean(values: list[float]) -> float:
    return float(sum(values) / len(values)) if values else float("nan")


def fmt_count(value: object, _: dict) -> str:
    if isinstance(value, str):
        return value
    return f"{int(float(value))}"


def fmt_percent(value: object, _: dict) -> str:
    return f"{100.0 * float(value):.1f}%"


def fmt_float3(value: object, _: dict) -> str:
    return f"{float(value):.3f}"


def success_rate(results: list[dict], threshold: float) -> float:
    flags = []
    for rec in results:
        add_s = float((rec.get("metrics") or {}).get("ADD_S", float("inf")))
        flags.append(1.0 if add_s <= threshold else 0.0)
    return mean(flags)


def grouped_by_obj(results: list[dict]) -> dict[int, list[dict]]:
    grouped: dict[int, list[dict]] = defaultdict(list)
    for rec in results:
        grouped[int(rec["obj_id"])].append(rec)
    return grouped


def build_rows(mapping: OrderedDict[int, str], baseline_results: list[dict], ours_results: list[dict], threshold: float) -> list[dict]:
    base_by_obj = grouped_by_obj(baseline_results)
    ours_by_obj = grouped_by_obj(ours_results)
    rows: list[dict] = []

    for obj_id, object_name in mapping.items():
        base_group = base_by_obj.get(obj_id, [])
        ours_group = ours_by_obj.get(obj_id, [])
        ours_metrics = [rec.get("metrics") or {} for rec in ours_group]
        ours_stats = [rec.get("stats") or {} for rec in ours_group]

        base_sr = success_rate(base_group, threshold)
        ours_sr = success_rate(ours_group, threshold)

        rows.append(
            {
                "Object": object_name,
                "N": len(ours_group),
                "Baseline SR": base_sr,
                "Ours SR": ours_sr,
                "Delta SR": ours_sr - base_sr,
                "Ours ADD": mean([float(m["ADD"]) for m in ours_metrics]),
                "Ours ADD-S": mean([float(m["ADD_S"]) for m in ours_metrics]),
                "Ours RE": mean([float(m["rotation_error_deg"]) for m in ours_metrics]),
                "Ours TE": mean([float(m["translation_error"]) for m in ours_metrics]),
                "Ours Time": mean([float(s["registration_time"]) for s in ours_stats]),
            }
        )

    if rows:
        n_objects = len(rows)
        rows.append(
            {
                "Object": "Mean",
                "N": f"{n_objects} objects",
                "Baseline SR": mean([float(r["Baseline SR"]) for r in rows]),
                "Ours SR": mean([float(r["Ours SR"]) for r in rows]),
                "Delta SR": mean([float(r["Delta SR"]) for r in rows]),
                "Ours ADD": mean([float(r["Ours ADD"]) for r in rows]),
                "Ours ADD-S": mean([float(r["Ours ADD-S"]) for r in rows]),
                "Ours RE": mean([float(r["Ours RE"]) for r in rows]),
                "Ours TE": mean([float(r["Ours TE"]) for r in rows]),
                "Ours Time": mean([float(r["Ours Time"]) for r in rows]),
            }
        )
    return rows


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--mapping_xlsx", required=True, type=Path)
    ap.add_argument("--baseline_json", required=True, type=Path)
    ap.add_argument("--ours_json", required=True, type=Path)
    ap.add_argument("--output_csv", required=True, type=Path)
    ap.add_argument("--threshold", type=float, default=20.0)
    args = ap.parse_args()

    mapping = load_mapping(args.mapping_xlsx)
    baseline_results = load_results(args.baseline_json)
    ours_results = load_results(args.ours_json)
    rows = build_rows(mapping, baseline_results, ours_results, args.threshold)
    export_table_bundle(
        args.output_csv.with_suffix(""),
        rows,
        [
            {**col, "format": fmt_count} if col["header"] == "N" else
            {**col, "format": fmt_percent} if col["header"] in {"Baseline SR", "Ours SR", "Delta SR"} else
            {**col, "format": fmt_float3} if col["align"] == "right" and col["header"] not in {"N", "Baseline SR", "Ours SR", "Delta SR"} else
            col
            for col in COLUMNS
        ],
        caption="Per-object LMO comparison between the same-backbone baseline and the full method. Success rate is computed with ADD-S <= 20 mm; distance metrics are reported in millimetres.",
        label="tab:lmo-per-object",
        bold_row_predicate=lambda row: row.get("Object") == "Mean",
        col_spec="@{}lrrrrrrrrr@{}",
    )
    print(f"saved {args.output_csv}")
    print(f"saved {args.output_csv.with_suffix('.md')}")
    print(f"saved {args.output_csv.with_suffix('.tex')}")
    print(f"rows {len(rows)}")


if __name__ == "__main__":
    main()
