from __future__ import annotations

import json
import math
import re
from collections import defaultdict
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Sequence

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.gridspec import GridSpec
import numpy as np
import open3d as o3d
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image
import yaml


ROOT = Path(__file__).resolve().parents[1]
STANFORD_RESULT_DIR = ROOT / "experiments" / "results" / "Stanford"
LMO_RESULT_DIR = ROOT / "experiments" / "results" / "lmo"
STANFORD_TABLE_DIR = ROOT / "experiments" / "tables" / "stanford"
STANFORD_ANALYSIS_DIR = STANFORD_TABLE_DIR / "analysis_plus"
LMO_TABLE_DIR = ROOT / "experiments" / "tables" / "lmo"

PAPER_DIR = ROOT / "paper_outputs"
TABLE_DIR = PAPER_DIR / "tables"
FIG_DIR = PAPER_DIR / "figures"
APP_TABLE_DIR = PAPER_DIR / "appendix_tables"
APP_FIG_DIR = PAPER_DIR / "appendix_figures"
ANALYSIS_DIR = PAPER_DIR / "analysis"

PNG_DPI = 320
STANFORD_SUCCESS_THRESHOLD = 0.02
LMO_PRIMARY_THRESHOLD_MM = 20.0

THIN = Side(style="thin", color="BFBFBF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
BEST_FILL = PatternFill("solid", fgColor="FFF2CC")
SECOND_FILL = PatternFill("solid", fgColor="E2F0D9")
ALT_FILL = PatternFill("solid", fgColor="F7F9FB")
HEADER_FONT = Font(name="Times New Roman", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Times New Roman", size=11)
TITLE_FONT = Font(name="Times New Roman", bold=True, size=13)

CORE_METHODS = [
    "Same-backbone Baseline",
    "No RS-MRQ",
    "No Robust Vote",
    "No Backend",
    "Ours (Full)",
]
TRADEOFF_METHODS = [
    "Adaptive Two-Stage",
    "Ours + RV20",
    "Ours + RV10",
    "Ours + RS-MRQ cap64",
    "Ours + RS-MRQ cap128",
    "Ours + RS-MRQ cap256",
]
EXTERNAL_METHODS = [
    "Drost",
    "Going Further",
    "Birdal Revisited",
    "Edge-enhanced PPF",
]
METHOD_ORDER = CORE_METHODS + TRADEOFF_METHODS + EXTERNAL_METHODS
METHOD_ORDER_INDEX = {name: i for i, name in enumerate(METHOD_ORDER)}
METHOD_GROUP = {
    **{name: "ablation" for name in CORE_METHODS},
    **{name: "tradeoff" for name in TRADEOFF_METHODS},
    **{name: "external" for name in EXTERNAL_METHODS},
}
METHOD_COLORS = {
    "Same-backbone Baseline": "#808080",
    "No RS-MRQ": "#4E79A7",
    "No Robust Vote": "#F28E2B",
    "No Backend": "#59A14F",
    "Ours (Full)": "#E15759",
    "Adaptive Two-Stage": "#76B7B2",
    "Ours + RV20": "#2F8F9D",
    "Ours + RV10": "#58B7C0",
    "Ours + RS-MRQ cap64": "#9C755F",
    "Ours + RS-MRQ cap128": "#C49A00",
    "Ours + RS-MRQ cap256": "#7F9E6E",
    "Drost": "#8C6D5A",
    "Going Further": "#6B6ECF",
    "Birdal Revisited": "#B5BD22",
    "Edge-enhanced PPF": "#17A2B8",
}
SEGMENT_COLORS = {
    "Scene preprocess": "#D0D0D0",
    "Front-end": "#F79024",
    "Back-end": "#76B7B2",
}

STANFORD_BATCHES = {
    "Same-backbone Baseline": STANFORD_RESULT_DIR / "stanford_same_backbone_baseline_batch.json",
    "No RS-MRQ": STANFORD_RESULT_DIR / "stanford_no_rsmrq_batch.json",
    "No Robust Vote": STANFORD_RESULT_DIR / "stanford_no_robust_vote_batch.json",
    "No Backend": STANFORD_RESULT_DIR / "stanford_no_backend_batch.json",
    "Ours (Full)": STANFORD_RESULT_DIR / "stanford_ours_full_batch.json",
    "Adaptive Two-Stage": STANFORD_RESULT_DIR / "stanford_adaptive_two_stage_batch.json",
    "Ours + RV20": STANFORD_RESULT_DIR / "stanford_ours_rv20_batch.json",
    "Ours + RV10": STANFORD_RESULT_DIR / "stanford_ours_rv10_batch.json",
    "Ours + RS-MRQ cap64": STANFORD_RESULT_DIR / "stanford_ours_cap64_batch.json",
    "Ours + RS-MRQ cap128": STANFORD_RESULT_DIR / "stanford_ours_cap128_batch.json",
    "Ours + RS-MRQ cap256": STANFORD_RESULT_DIR / "stanford_ours_cap256_batch.json",
    "Drost": STANFORD_RESULT_DIR / "stanford_drost_original_batch.json",
    "Going Further": STANFORD_RESULT_DIR / "stanford_going_further_ppf_batch.json",
    "Birdal Revisited": STANFORD_RESULT_DIR / "stanford_birdal_revisited_batch.json",
    "Edge-enhanced PPF": STANFORD_RESULT_DIR / "stanford_edge_enhanced_ppf_batch.json",
}
LMO_BATCHES = {
    "Same-backbone Baseline": LMO_RESULT_DIR / "lmo_same_backbone_batch.json",
    "No RS-MRQ": LMO_RESULT_DIR / "lmo_no_rsmrq_batch.json",
    "No Robust Vote": LMO_RESULT_DIR / "lmo_no_robust_vote_batch.json",
    "No Backend": LMO_RESULT_DIR / "lmo_no_backend_batch.json",
    "Ours (Full)": LMO_RESULT_DIR / "lmo_ours_full_batch.json",
    "Adaptive Two-Stage": LMO_RESULT_DIR / "lmo_adaptive_batch.json",
    "Ours + RV20": LMO_RESULT_DIR / "lmo_ours_rv20_batch.json",
    "Ours + RV10": LMO_RESULT_DIR / "lmo_ours_rv10_batch.json",
    "Ours + RS-MRQ cap64": LMO_RESULT_DIR / "lmo_ours_cap64_batch.json",
    "Ours + RS-MRQ cap128": LMO_RESULT_DIR / "lmo_ours_cap128_batch.json",
    "Ours + RS-MRQ cap256": LMO_RESULT_DIR / "lmo_ours_cap256_batch.json",
    "Drost": LMO_RESULT_DIR / "lmo_drost_batch.json",
    "Going Further": LMO_RESULT_DIR / "lmo_going_further_batch.json",
    "Birdal Revisited": LMO_RESULT_DIR / "lmo_birdal_batch.json",
    "Edge-enhanced PPF": LMO_RESULT_DIR / "lmo_edge_enhanced_ppf_batch.json",
}

MODULE_FLAGS = {
    "Same-backbone Baseline": {
        "RS-MRQ": "No",
        "Robust Vote": "No",
        "Candidate-level Selection": "No",
        "Mode Clustering": "No",
    },
    "No RS-MRQ": {
        "RS-MRQ": "No",
        "Robust Vote": "Yes",
        "Candidate-level Selection": "Yes",
        "Mode Clustering": "Yes",
    },
    "No Robust Vote": {
        "RS-MRQ": "Yes",
        "Robust Vote": "No",
        "Candidate-level Selection": "Yes",
        "Mode Clustering": "Yes",
    },
    "No Backend": {
        "RS-MRQ": "Yes",
        "Robust Vote": "Yes",
        "Candidate-level Selection": "No",
        "Mode Clustering": "No",
    },
    "Ours (Full)": {
        "RS-MRQ": "Yes",
        "Robust Vote": "Yes",
        "Candidate-level Selection": "Yes",
        "Mode Clustering": "Yes",
    },
}


def ensure_dirs() -> None:
    for path in [PAPER_DIR, TABLE_DIR, FIG_DIR, APP_TABLE_DIR, APP_FIG_DIR, ANALYSIS_DIR]:
        path.mkdir(parents=True, exist_ok=True)


def apply_plot_style() -> None:
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
            "mathtext.fontset": "stix",
            "axes.titlesize": 12,
            "axes.labelsize": 11,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
            "legend.fontsize": 9,
            "axes.grid": True,
            "grid.linestyle": "--",
            "grid.alpha": 0.3,
            "axes.linewidth": 0.8,
            "savefig.bbox": "tight",
        }
    )


def safe_float(value: Any) -> float:
    try:
        out = float(value)
    except (TypeError, ValueError):
        return float("nan")
    if not np.isfinite(out):
        return float("nan")
    return out


def nanmean(values: Iterable[Any]) -> float:
    arr = np.asarray([safe_float(v) for v in values], dtype=float)
    arr = arr[np.isfinite(arr)]
    return float(arr.mean()) if arr.size else float("nan")


def nanmedian(values: Iterable[Any]) -> float:
    arr = np.asarray([safe_float(v) for v in values], dtype=float)
    arr = arr[np.isfinite(arr)]
    return float(np.median(arr)) if arr.size else float("nan")


def nanmax(values: Iterable[Any]) -> float:
    arr = np.asarray([safe_float(v) for v in values], dtype=float)
    arr = arr[np.isfinite(arr)]
    return float(arr.max()) if arr.size else float("nan")


def success_rate(values: Iterable[Any], threshold: float) -> float:
    arr = np.asarray([safe_float(v) for v in values], dtype=float)
    arr = arr[np.isfinite(arr)]
    return float((arr <= threshold).mean()) if arr.size else float("nan")


def infer_case_key(dataset: str, rec: Mapping[str, Any]) -> str:
    if dataset == "Stanford":
        return f"{rec.get('scene_name')}|{rec.get('scene_variant')}|{rec.get('obj_id')}"
    row = rec.get("row", {}) or {}
    frame_id = row.get("frame_id", rec.get("scene_variant"))
    scene_id = row.get("scene_id", rec.get("scene_id"))
    return f"{scene_id}|{frame_id}|{rec.get('obj_id')}"


def pretty_scene(scene_name: str) -> str:
    if scene_name.startswith("LMO_"):
        return scene_name.replace("LMO_", "LMO Scene ")
    return scene_name


def resolve_repo_path(path_like: Any) -> Path | None:
    if not path_like:
        return None
    text = str(path_like)
    path = Path(text)
    if path.exists():
        return path
    normalized = text.replace("\\", "/")
    for anchor in ("configs/", "data/", "experiments/"):
        if anchor in normalized:
            suffix = normalized[normalized.index(anchor) :]
            candidate = ROOT / Path(suffix)
            if candidate.exists():
                return candidate
    if not path.is_absolute():
        candidate = ROOT / path
        if candidate.exists():
            return candidate
    return path


@lru_cache(maxsize=None)
def load_yaml_config(path_like: str) -> dict[str, Any]:
    path = resolve_repo_path(path_like)
    if path is None or not path.exists():
        return {}
    try:
        return yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except Exception:
        return {}


def nested_get(obj: Mapping[str, Any], *keys: str, default: Any = float("nan")) -> Any:
    cur: Any = obj
    for key in keys:
        if not isinstance(cur, Mapping) or key not in cur:
            return default
        cur = cur[key]
    return cur


def compute_top_score_margin(scores: Sequence[Any] | None) -> float:
    if not scores or len(scores) < 2:
        return float("nan")
    return safe_float(scores[0]) - safe_float(scores[1])


def resolve_scene_and_model_paths(dataset: str, rec: Mapping[str, Any]) -> tuple[Path | None, Path | None]:
    if dataset == "LMO":
        row = rec.get("row", {}) or {}
        scene_path = resolve_repo_path(row.get("pcd_path") or rec.get("scene_path"))
        model_path = resolve_repo_path(row.get("expected_model_path") or rec.get("model_path"))
        return scene_path, model_path
    return resolve_repo_path(rec.get("scene_path")), resolve_repo_path(rec.get("model_path"))


def load_batch_records(dataset: str, method: str, batch_path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    payload = json.loads(batch_path.read_text(encoding="utf-8"))
    config_path = str(payload.get("config_path") or "")
    cfg = load_yaml_config(config_path)
    rv_top_m = nested_get(cfg, "robust_vote", "top_m_per_bucket", default=float("nan"))
    rsmrq_cap = nested_get(cfg, "rsmrq", "global_candidate_cap", default=float("nan"))
    method_rows: list[dict[str, Any]] = []
    render_rows: list[dict[str, Any]] = []
    for rec in payload.get("results", []) or []:
        stats = rec.get("stats", {}) or {}
        metrics = rec.get("metrics", {}) or {}
        debug = rec.get("debug", {}) or {}
        row = rec.get("row", {}) or {}
        pose_sel = debug.get("pose_selection", {}) or {}
        rsmrq = debug.get("rsmrq", {}) or {}
        robust_vote = debug.get("robust_vote", {}) or {}
        case_key = infer_case_key(dataset, rec)
        scene_path, model_path = resolve_scene_and_model_paths(dataset, rec)
        top_scores = pose_sel.get("top_scores") or []
        method_rows.append(
            {
                "dataset": dataset,
                "method": method,
                "method_group": METHOD_GROUP.get(method, "other"),
                "method_order": METHOD_ORDER_INDEX.get(method, 999),
                "source_batch_json": str(batch_path),
                "config_path": config_path,
                "idx": rec.get("idx"),
                "case_id": case_key,
                "scene_id": rec.get("scene_id", row.get("scene_id")),
                "scene_name": rec.get("scene_name", row.get("scene_name")),
                "scene_variant": rec.get("scene_variant", row.get("scene_variant")),
                "obj_id": rec.get("obj_id", row.get("obj_id")),
                "frame_id": row.get("frame_id", rec.get("scene_variant")),
                "ADD": safe_float(metrics.get("ADD")),
                "ADD_S": safe_float(metrics.get("ADD_S")),
                "rotation_error_deg": safe_float(metrics.get("rotation_error_deg")),
                "translation_error": safe_float(metrics.get("translation_error")),
                "inlier_ratio": safe_float(metrics.get("inlier_ratio")),
                "registration_time": safe_float(stats.get("registration_time")),
                "total_time": safe_float(stats.get("total_time")),
                "scene_preprocess_time": safe_float(stats.get("scene_preprocess_time")),
                "ppf_frontend_time": safe_float(stats.get("ppf_frontend_time")),
                "backend_time": safe_float(stats.get("backend_time")),
                "pose_selection_time": safe_float(stats.get("pose_selection_time")),
                "pose_clustering_time": safe_float(stats.get("pose_clustering_time")),
                "legacy_clustering_time": safe_float(stats.get("legacy_clustering_time")),
                "candidate_inflation_mean": safe_float(
                    stats.get("candidate_inflation_mean", debug.get("candidate_inflation_mean"))
                ),
                "final_pose_policy": stats.get("final_pose_policy", debug.get("final_pose_policy")),
                "adaptive_used_upgrade": rec.get("adaptive_used_upgrade"),
                "adaptive_trigger_reason": rec.get("adaptive_trigger_reason"),
                "adaptive_total_effective_time": safe_float(rec.get("adaptive_total_effective_time")),
                "adaptive_stage1_registration_time": safe_float(rec.get("adaptive_stage1_registration_time")),
                "adaptive_stage2_registration_time": safe_float(rec.get("adaptive_stage2_registration_time")),
                "adaptive_stage1_frontend_time": safe_float(rec.get("adaptive_stage1_frontend_time")),
                "adaptive_stage2_frontend_time": safe_float(rec.get("adaptive_stage2_frontend_time")),
                "adaptive_stage1_method": rec.get("adaptive_stage1_method"),
                "adaptive_stage2_method": rec.get("adaptive_stage2_method"),
                "adaptive_final_source": rec.get("adaptive_final_source"),
                "num_input_candidates": safe_float(pose_sel.get("num_input_candidates")),
                "num_selected_candidates": safe_float(pose_sel.get("num_selected")),
                "num_preselected_candidates": safe_float(pose_sel.get("num_preselected")),
                "best_score": safe_float(pose_sel.get("best_score")),
                "best_visibility_support": safe_float(pose_sel.get("best_visibility_support")),
                "top_scores_margin": compute_top_score_margin(top_scores),
                "n_votes": safe_float(robust_vote.get("n_votes")),
                "topm_used": safe_float(robust_vote.get("topm_used")),
                "topm_total": safe_float(robust_vote.get("topm_total")),
                "topm_ratio": safe_float(robust_vote.get("topm_ratio")),
                "rsmrq_global_candidate_cap": safe_float(rsmrq.get("global_candidate_cap", rsmrq_cap)),
                "rsmrq_cap_hit_count": safe_float(rsmrq.get("cap_hit_count")),
                "rsmrq_cap_hit_ratio": safe_float(rsmrq.get("cap_hit_ratio")),
                "rsmrq_merged_candidates_pre_cap_mean": safe_float(rsmrq.get("merged_candidates_pre_cap_mean")),
                "rsmrq_merged_candidates_post_cap_mean": safe_float(rsmrq.get("merged_candidates_post_cap_mean")),
                "robust_vote_top_m_per_bucket": safe_float(rv_top_m),
                "robust_vote_enabled": bool(nested_get(cfg, "robust_vote", "enable", default=True)),
                "rsmrq_enabled": bool(nested_get(cfg, "rsmrq", "enable", default=True)),
                "success_add_002": safe_float(metrics.get("ADD")) <= STANFORD_SUCCESS_THRESHOLD,
                "success_adds_002": safe_float(metrics.get("ADD_S")) <= STANFORD_SUCCESS_THRESHOLD,
                "success_rot_5deg": safe_float(metrics.get("rotation_error_deg")) <= 5.0,
                "primary_success": (
                    safe_float(metrics.get("ADD")) <= STANFORD_SUCCESS_THRESHOLD
                    if dataset == "Stanford"
                    else safe_float(metrics.get("ADD_S")) <= LMO_PRIMARY_THRESHOLD_MM
                ),
                "primary_accuracy_name": (
                    "SR@ADD<=0.02" if dataset == "Stanford" else "SR@ADD-S<=20mm"
                ),
            }
        )
        render_rows.append(
            {
                "dataset": dataset,
                "method": method,
                "case_id": case_key,
                "scene_name": rec.get("scene_name", row.get("scene_name")),
                "scene_variant": rec.get("scene_variant", row.get("scene_variant")),
                "obj_id": rec.get("obj_id", row.get("obj_id")),
                "frame_id": row.get("frame_id", rec.get("scene_variant")),
                "scene_path": scene_path,
                "model_path": model_path,
                "T_pred": np.asarray(rec.get("T_pred"), dtype=np.float64),
                "T_gt": np.asarray(rec.get("T_gt"), dtype=np.float64),
                "ADD": safe_float(metrics.get("ADD")),
                "ADD_S": safe_float(metrics.get("ADD_S")),
                "registration_time": safe_float(stats.get("registration_time")),
            }
        )
    return method_rows, render_rows


def build_master_tables() -> tuple[pd.DataFrame, pd.DataFrame, dict[tuple[str, str, str], dict[str, Any]]]:
    rows: list[dict[str, Any]] = []
    render_lookup: dict[tuple[str, str, str], dict[str, Any]] = {}
    for dataset, mapping in [("Stanford", STANFORD_BATCHES), ("LMO", LMO_BATCHES)]:
        for method, batch_path in mapping.items():
            if not batch_path.exists():
                continue
            method_rows, render_rows = load_batch_records(dataset, method, batch_path)
            rows.extend(method_rows)
            for item in render_rows:
                render_lookup[(dataset, method, item["case_id"])] = item
    master = pd.DataFrame(rows)
    master = master.sort_values(["dataset", "method_order", "idx", "obj_id"], kind="stable").reset_index(drop=True)

    summary_rows: list[dict[str, Any]] = []
    for (dataset, method), gdf in master.groupby(["dataset", "method"], sort=False):
        add_vals = gdf["ADD"]
        adds_vals = gdf["ADD_S"]
        reg_vals = gdf["registration_time"]
        front_vals = gdf["ppf_frontend_time"]
        back_vals = gdf["backend_time"]
        summary_rows.append(
            {
                "dataset": dataset,
                "method": method,
                "method_group": METHOD_GROUP.get(method, "other"),
                "method_order": METHOD_ORDER_INDEX.get(method, 999),
                "n_cases": int(len(gdf)),
                "mean_registration_time": nanmean(reg_vals),
                "median_registration_time": nanmedian(reg_vals),
                "max_registration_time": nanmax(reg_vals),
                "mean_frontend_time": nanmean(front_vals),
                "mean_backend_time": nanmean(back_vals),
                "mean_scene_preprocess_time": nanmean(gdf["scene_preprocess_time"]),
                "mean_ADD": nanmean(add_vals),
                "mean_ADD_S": nanmean(adds_vals),
                "mean_rotation_error_deg": nanmean(gdf["rotation_error_deg"]),
                "mean_translation_error": nanmean(gdf["translation_error"]),
                "SR@ADD<=0.02": success_rate(add_vals, STANFORD_SUCCESS_THRESHOLD),
                "SR@ADD-S<=0.02": success_rate(adds_vals, STANFORD_SUCCESS_THRESHOLD),
                "SR@rot<=5deg": float(pd.Series(gdf["success_rot_5deg"]).mean()) if len(gdf) else float("nan"),
                "candidate_inflation_mean": nanmean(gdf["candidate_inflation_mean"]),
                "num_input_candidates_mean": nanmean(gdf["num_input_candidates"]),
                "n_votes_mean": nanmean(gdf["n_votes"]),
                "topm_ratio_mean": nanmean(gdf["topm_ratio"]),
                "primary_accuracy_name": "SR@ADD<=0.02" if dataset == "Stanford" else "SR@ADD-S<=20mm",
                "primary_accuracy": (
                    float(pd.Series(gdf["primary_success"]).mean()) if len(gdf) else float("nan")
                ),
            }
        )
    summary = pd.DataFrame(summary_rows).sort_values(
        ["dataset", "method_order"], kind="stable"
    ).reset_index(drop=True)
    return master, summary, render_lookup


def style_worksheet(ws, freeze_cell: str | None = "A2") -> None:
    if freeze_cell:
        ws.freeze_panes = freeze_cell
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    sample_rows = min(ws.max_row, 200)
    for col_idx in range(1, ws.max_column + 1):
        values = [
            str(ws.cell(row=r, column=col_idx).value or "")
            for r in range(1, sample_rows + 1)
        ]
        width = min(max(len(v) for v in values) + 2, 36)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def df_to_workbook(path: Path, sheets: Mapping[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        style_worksheet(ws)
    wb.save(path)


def apply_sheet_number_formats(path: Path, sheet_to_formats: Mapping[str, Mapping[str, str]]) -> None:
    wb = load_workbook(path)
    for sheet_name, col_formats in sheet_to_formats.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        for header, number_format in col_formats.items():
            col_idx = headers.get(header)
            if col_idx is None:
                continue
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row, col_idx)
                if isinstance(cell.value, (int, float, np.floating)):
                    cell.number_format = number_format
    wb.save(path)


def write_block(
    ws,
    title: str,
    df: pd.DataFrame,
    start_row: int,
    highlight_rules: Mapping[str, str] | None = None,
    number_formats: Mapping[str, str] | None = None,
) -> int:
    ws.cell(start_row, 1, title)
    ws.cell(start_row, 1).font = TITLE_FONT
    ws.cell(start_row, 1).fill = SECTION_FILL
    start_row += 1
    headers = list(df.columns)
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(start_row, col_idx, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    data_start = start_row + 1
    for row_offset, (_, row) in enumerate(df.iterrows(), start=0):
        excel_row = data_start + row_offset
        for col_idx, header in enumerate(headers, start=1):
            val = row[header]
            cell = ws.cell(excel_row, col_idx, None if pd.isna(val) else val)
            cell.font = BODY_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if row_offset % 2 == 1:
                cell.fill = ALT_FILL
            if number_formats and header in number_formats and isinstance(val, (int, float, np.floating)):
                cell.number_format = number_formats[header]
    if highlight_rules:
        for header, direction in highlight_rules.items():
            if header not in headers:
                continue
            series = pd.to_numeric(df[header], errors="coerce")
            valid = series.dropna()
            if valid.empty:
                continue
            sorted_index = (
                valid.sort_values(ascending=(direction == "min")).index.tolist()
            )
            best_idx = sorted_index[0]
            best_cell = ws.cell(data_start + df.index.get_loc(best_idx), headers.index(header) + 1)
            best_cell.fill = BEST_FILL
            best_cell.font = Font(name="Times New Roman", bold=True)
            if len(sorted_index) > 1:
                second_idx = sorted_index[1]
                second_cell = ws.cell(data_start + df.index.get_loc(second_idx), headers.index(header) + 1)
                second_cell.fill = SECOND_FILL
                second_cell.font = Font(name="Times New Roman", underline="single")
    return data_start + len(df)


def finalize_workbook(path: Path, sheet_name: str = "Sheet1") -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for col in ws.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in col]
            width = min(max(len(v) for v in values) + 2, 40)
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
    wb.save(path)


def write_specialized_workbook(path: Path, blocks: Sequence[tuple[str, pd.DataFrame, Mapping[str, str] | None]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    row = 1
    for title, df, rules in blocks:
        row = write_block(
            ws,
            title,
            df,
            row,
            highlight_rules=rules,
            number_formats={
                c: "0.0000"
                for c in df.columns
                if c not in {"Method", "Variant", "Dataset", "RS-MRQ", "Robust Vote", "Candidate-level Selection", "Mode Clustering", "Final Choice", "Note", "Primary Accuracy Metric"}
            },
        )
        row += 2
    wb.save(path)
    finalize_workbook(path)


def save_figure(fig: plt.Figure, out_base: Path) -> None:
    fig.savefig(out_base.with_suffix(".png"), dpi=PNG_DPI, bbox_inches="tight")
    fig.savefig(out_base.with_suffix(".svg"), bbox_inches="tight")
    plt.close(fig)


def add_panel_label(ax, label: str) -> None:
    ax.text(
        0.01,
        0.99,
        label,
        transform=ax.transAxes,
        ha="left",
        va="top",
        fontsize=11,
        fontweight="bold",
        bbox=dict(facecolor="white", edgecolor="none", alpha=0.85, pad=0.2),
    )


def add_subfigure_labels_below(
    fig: plt.Figure,
    axes: Sequence[plt.Axes],
    labels: Sequence[str],
    offset: float = 0.08,
    fixed_y: float | None = None,
) -> None:
    for ax, label in zip(axes, labels, strict=True):
        bbox = ax.get_position()
        x = (bbox.x0 + bbox.x1) / 2.0
        y = fixed_y if fixed_y is not None else max(0.01, bbox.y0 - offset)
        fig.text(x, y, label, ha="center", va="top", fontsize=10.5, fontweight="bold")


def axis_limits(values: np.ndarray, lower: float | None = None, upper: float | None = None) -> tuple[float, float]:
    finite = values[np.isfinite(values)]
    if finite.size == 0:
        return 0.0, 1.0
    lo = float(finite.min())
    hi = float(finite.max())
    span = max(hi - lo, 1e-6)
    pad = span * 0.12
    out_lo = lo - pad
    out_hi = hi + pad
    if lower is not None:
        out_lo = max(lower, out_lo)
    if upper is not None:
        out_hi = min(upper, out_hi)
    return out_lo, out_hi


def summary_lookup(summary_df: pd.DataFrame, dataset: str) -> pd.DataFrame:
    return summary_df[summary_df["dataset"] == dataset].set_index("method")


def build_master_outputs(master: pd.DataFrame, summary: pd.DataFrame) -> None:
    df_to_workbook(ANALYSIS_DIR / "master_per_case.xlsx", {"PerCase": master})
    df_to_workbook(ANALYSIS_DIR / "master_summary.xlsx", {"Summary": summary})


def build_main_results_table(summary: pd.DataFrame, out_path: Path | None = None) -> None:
    stanford_methods = [
        "Same-backbone Baseline",
        "No RS-MRQ",
        "No Robust Vote",
        "No Backend",
        "Ours (Full)",
        "Adaptive Two-Stage",
        "Drost",
        "Going Further",
        "Birdal Revisited",
        "Edge-enhanced PPF",
    ]
    lmo_methods = [
        "Same-backbone Baseline",
        "No RS-MRQ",
        "No Robust Vote",
        "No Backend",
        "Ours (Full)",
        "Adaptive Two-Stage",
        "Drost",
        "Going Further",
        "Birdal Revisited",
        "Edge-enhanced PPF",
    ]
    s = summary_lookup(summary, "Stanford").loc[stanford_methods].reset_index()
    l = summary_lookup(summary, "LMO").loc[lmo_methods].reset_index()
    stanford_block = pd.DataFrame(
        {
            "Method": s["method"],
            "ADD": s["mean_ADD"],
            "ADD-S": s["mean_ADD_S"],
            "SR@ADD<=0.02": s["SR@ADD<=0.02"],
            "Rot. Err.": s["mean_rotation_error_deg"],
            "Trans. Err.": s["mean_translation_error"],
            "Reg. Time": s["mean_registration_time"],
        }
    )
    lmo_block = pd.DataFrame(
        {
            "Method": l["method"],
            "ADD": l["mean_ADD"],
            "ADD-S": l["mean_ADD_S"],
            "SR@ADD-S<=20mm": l["primary_accuracy"],
            "Rot. Err.": l["mean_rotation_error_deg"],
            "Trans. Err.": l["mean_translation_error"],
            "Reg. Time": l["mean_registration_time"],
        }
    )
    write_specialized_workbook(
        out_path or (TABLE_DIR / "main_results_stanford_lmo.xlsx"),
        [
            (
                "Block A: Stanford",
                stanford_block,
                {
                    "ADD": "min",
                    "ADD-S": "min",
                    "SR@ADD<=0.02": "max",
                    "Rot. Err.": "min",
                    "Trans. Err.": "min",
                    "Reg. Time": "min",
                },
            ),
            (
                "Block B: LMO",
                lmo_block,
                {
                    "ADD": "min",
                    "ADD-S": "min",
                    "SR@ADD-S<=20mm": "max",
                    "Rot. Err.": "min",
                    "Trans. Err.": "min",
                    "Reg. Time": "min",
                },
            ),
        ],
    )


def build_overall_ablation_table(summary: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stanford_Ablation"
    s = summary_lookup(summary, "Stanford").loc[CORE_METHODS].reset_index()
    l = summary_lookup(summary, "LMO").loc[CORE_METHODS].reset_index()
    stanford_block = []
    for _, row in s.iterrows():
        stanford_block.append(
            {
                "Method": row["method"],
                **MODULE_FLAGS[row["method"]],
                "ADD": row["mean_ADD"],
                "SR@ADD<=0.02": row["SR@ADD<=0.02"],
                "Reg. Time": row["mean_registration_time"],
                "Front-end Time": row["mean_frontend_time"],
                "Back-end Time": row["mean_backend_time"],
            }
        )
    lmo_block = []
    for _, row in l.iterrows():
        lmo_block.append(
            {
                "Method": row["method"],
                **MODULE_FLAGS[row["method"]],
                "Primary Accuracy Metric": row["primary_accuracy"],
                "Reg. Time": row["mean_registration_time"],
            }
        )
    write_block(
        ws,
        "Stanford Ablation",
        pd.DataFrame(stanford_block),
        1,
        highlight_rules={
            "ADD": "min",
            "SR@ADD<=0.02": "max",
            "Reg. Time": "min",
            "Front-end Time": "min",
            "Back-end Time": "min",
        },
    )
    ws2 = wb.create_sheet("LMO_Ablation")
    write_block(
        ws2,
        "LMO Ablation",
        pd.DataFrame(lmo_block),
        1,
        highlight_rules={"Primary Accuracy Metric": "max", "Reg. Time": "min"},
    )
    wb.save(TABLE_DIR / "overall_ablation.xlsx")
    finalize_workbook(TABLE_DIR / "overall_ablation.xlsx")


def build_efficiency_extensions_table(summary: pd.DataFrame) -> None:
    ours_s = float(summary_lookup(summary, "Stanford").loc["Ours (Full)", "mean_registration_time"])
    ours_s_acc = float(summary_lookup(summary, "Stanford").loc["Ours (Full)", "SR@ADD<=0.02"])
    selected = ["No RS-MRQ", "Ours (Full)", "Adaptive Two-Stage", "Ours + RV20", "Ours + RV10", "Ours + RS-MRQ cap256"]
    s = summary_lookup(summary, "Stanford").loc[selected].reset_index()
    s_rows = []
    for _, row in s.iterrows():
        note = "-"
        if row["method"] == "Adaptive Two-Stage":
            note = "Upgrade Ratio from adaptive batch"
        elif "RV" in row["method"]:
            note = row["method"].replace("Ours + ", "")
        elif "cap" in row["method"]:
            note = row["method"].replace("Ours + ", "")
        s_rows.append(
            {
                "Method": row["method"],
                "Mean ADD": row["mean_ADD"],
                "SR@ADD<=0.02": row["SR@ADD<=0.02"],
                "Mean Registration Time": row["mean_registration_time"],
                "Mean Front-end Time": row["mean_frontend_time"],
                "Mean Back-end Time": row["mean_backend_time"],
                "Relative Time vs Ours": (row["mean_registration_time"] / ours_s - 1.0) * 100.0,
                "Relative Accuracy vs Ours": row["SR@ADD<=0.02"] - ours_s_acc,
                "Note": note,
            }
        )
    out_path = TABLE_DIR / "efficiency_extensions.xlsx"
    df_to_workbook(out_path, {"Stanford_Efficiency": pd.DataFrame(s_rows)})
    apply_sheet_number_formats(
        out_path,
        {
            "Stanford_Efficiency": {
                "Mean ADD": "0.000",
                "SR@ADD<=0.02": "0.00%",
                "Mean Registration Time": "0.000",
                "Mean Front-end Time": "0.000",
                "Mean Back-end Time": "0.000",
                "Relative Time vs Ours": "0.000",
                "Relative Accuracy vs Ours": "0.000",
            }
        },
    )


def build_parameter_selection_table(summary: pd.DataFrame) -> None:
    s = summary_lookup(summary, "Stanford")
    rows = [
        {
            "Variant": "RS-MRQ original",
            "SR@ADD<=0.02": s.loc["Ours (Full)", "SR@ADD<=0.02"],
            "Reg. Time": s.loc["Ours (Full)", "mean_registration_time"],
            "Front-end Time": s.loc["Ours (Full)", "mean_frontend_time"],
            "Relative Gain / Cost": "Reference",
            "Final Choice": "No",
            "Note": "No cap",
        },
        {
            "Variant": "RS-MRQ cap256",
            "SR@ADD<=0.02": s.loc["Ours + RS-MRQ cap256", "SR@ADD<=0.02"],
            "Reg. Time": s.loc["Ours + RS-MRQ cap256", "mean_registration_time"],
            "Front-end Time": s.loc["Ours + RS-MRQ cap256", "mean_frontend_time"],
            "Relative Gain / Cost": f"{(s.loc['Ours + RS-MRQ cap256', 'mean_registration_time'] / s.loc['Ours (Full)', 'mean_registration_time'] - 1.0) * 100.0:+.2f}% time",
            "Final Choice": "Yes",
            "Note": "Best cap under equal Stanford accuracy",
        },
        {
            "Variant": "RS-MRQ cap128",
            "SR@ADD<=0.02": s.loc["Ours + RS-MRQ cap128", "SR@ADD<=0.02"],
            "Reg. Time": s.loc["Ours + RS-MRQ cap128", "mean_registration_time"],
            "Front-end Time": s.loc["Ours + RS-MRQ cap128", "mean_frontend_time"],
            "Relative Gain / Cost": f"{(s.loc['Ours + RS-MRQ cap128', 'mean_registration_time'] / s.loc['Ours (Full)', 'mean_registration_time'] - 1.0) * 100.0:+.2f}% time",
            "Final Choice": "No",
            "Note": "Cap barely triggers",
        },
        {
            "Variant": "RS-MRQ cap64",
            "SR@ADD<=0.02": s.loc["Ours + RS-MRQ cap64", "SR@ADD<=0.02"],
            "Reg. Time": s.loc["Ours + RS-MRQ cap64", "mean_registration_time"],
            "Front-end Time": s.loc["Ours + RS-MRQ cap64", "mean_frontend_time"],
            "Relative Gain / Cost": f"{(s.loc['Ours + RS-MRQ cap64', 'mean_registration_time'] / s.loc['Ours (Full)', 'mean_registration_time'] - 1.0) * 100.0:+.2f}% time",
            "Final Choice": "No",
            "Note": "Triggers but hurts runtime",
        },
        {
            "Variant": "Robust Vote original",
            "SR@ADD<=0.02": s.loc["Ours (Full)", "SR@ADD<=0.02"],
            "Reg. Time": s.loc["Ours (Full)", "mean_registration_time"],
            "Front-end Time": s.loc["Ours (Full)", "mean_frontend_time"],
            "Relative Gain / Cost": "Reference",
            "Final Choice": "No",
            "Note": "top-m=50",
        },
        {
            "Variant": "Robust Vote rv20",
            "SR@ADD<=0.02": s.loc["Ours + RV20", "SR@ADD<=0.02"],
            "Reg. Time": s.loc["Ours + RV20", "mean_registration_time"],
            "Front-end Time": s.loc["Ours + RV20", "mean_frontend_time"],
            "Relative Gain / Cost": f"{(s.loc['Ours + RV20', 'mean_registration_time'] / s.loc['Ours (Full)', 'mean_registration_time'] - 1.0) * 100.0:+.2f}% time",
            "Final Choice": "Yes",
            "Note": "Best safe RV light point",
        },
        {
            "Variant": "Robust Vote rv10",
            "SR@ADD<=0.02": s.loc["Ours + RV10", "SR@ADD<=0.02"],
            "Reg. Time": s.loc["Ours + RV10", "mean_registration_time"],
            "Front-end Time": s.loc["Ours + RV10", "mean_frontend_time"],
            "Relative Gain / Cost": f"{(s.loc['Ours + RV10', 'mean_registration_time'] / s.loc['Ours (Full)', 'mean_registration_time'] - 1.0) * 100.0:+.2f}% time",
            "Final Choice": "Optional",
            "Note": "Aggressive version",
        },
        {
            "Variant": "Adaptive final trigger",
            "SR@ADD<=0.02": s.loc["Adaptive Two-Stage", "SR@ADD<=0.02"],
            "Reg. Time": s.loc["Adaptive Two-Stage", "mean_registration_time"],
            "Front-end Time": s.loc["Adaptive Two-Stage", "mean_frontend_time"],
            "Relative Gain / Cost": f"{(s.loc['Adaptive Two-Stage', 'mean_registration_time'] / s.loc['Ours (Full)', 'mean_registration_time'] - 1.0) * 100.0:+.2f}% time",
            "Final Choice": "Yes",
            "Note": "best_score<0.72 or margin<0.02 or visibility<0.28",
        },
    ]
    out_path = TABLE_DIR / "parameter_selection_summary.xlsx"
    df_to_workbook(out_path, {"ParameterSelection": pd.DataFrame(rows)})
    apply_sheet_number_formats(
        out_path,
        {
            "ParameterSelection": {
                "SR@ADD<=0.02": "0.00%",
                "Reg. Time": "0.000",
                "Front-end Time": "0.000",
            }
        },
    )


def plot_main_ablation_figure(summary: pd.DataFrame) -> None:
    s = summary_lookup(summary, "Stanford").loc[CORE_METHODS].reset_index()
    fig, axes = plt.subplots(1, 2, figsize=(7.1, 2.95), gridspec_kw={"width_ratios": [1.0, 1.06]})
    ax = axes[0]
    x = s["mean_registration_time"].to_numpy(dtype=float)
    y = s["SR@ADD<=0.02"].to_numpy(dtype=float)
    ax.set_xlim(*axis_limits(x, lower=0.0))
    ax.set_ylim(*axis_limits(y, lower=0.0, upper=1.0))
    label_text = {
        "Same-backbone Baseline": "Same-backbone",
        "No RS-MRQ": "No RS-MRQ",
        "No Robust Vote": "No RV",
        "No Backend": "No Backend",
        "Ours (Full)": "Ours",
    }
    label_positions = {
        "Same-backbone Baseline": {"dx": 0.45, "dy": 0.008, "ha": "left"},
        "No RS-MRQ": {"dx": -0.32, "dy": 0.010, "ha": "right"},
        "No Robust Vote": {"dx": 0.45, "dy": 0.008, "ha": "left"},
        "No Backend": {"dx": 0.62, "dy": -0.0015, "ha": "left"},
        "Ours (Full)": {"dx": 0.62, "dy": 0.013, "ha": "left"},
    }
    for _, row in s.iterrows():
        method = row["method"]
        is_ours = method == "Ours (Full)"
        ax.scatter(
            row["mean_registration_time"],
            row["SR@ADD<=0.02"],
            s=150 if is_ours else 110,
            color=METHOD_COLORS[method],
            edgecolor="black" if is_ours else "white",
            linewidth=1.2 if is_ours else 0.8,
            zorder=3,
        )
        pos = label_positions[method]
        tx = row["mean_registration_time"] + pos["dx"]
        ty = row["SR@ADD<=0.02"] + pos["dy"]
        ax.annotate(
            label_text[method],
            xy=(row["mean_registration_time"], row["SR@ADD<=0.02"]),
            xytext=(tx, ty),
            textcoords="data",
            ha=pos["ha"],
            va="center",
            fontsize=9,
            bbox=dict(facecolor="white", edgecolor="none", alpha=0.78, pad=0.15),
        )
    ax.set_title("Accuracy vs. Runtime")
    ax.set_xlabel("Mean Registration Time (s)")
    ax.set_ylabel("SR@ADD<=0.02")

    ax2 = axes[1]
    xpos = np.arange(len(CORE_METHODS))
    preprocess = s["mean_scene_preprocess_time"].to_numpy(dtype=float)
    frontend = s["mean_frontend_time"].to_numpy(dtype=float)
    backend = s["mean_backend_time"].to_numpy(dtype=float)
    ax2.bar(xpos, preprocess, color=SEGMENT_COLORS["Scene preprocess"], width=0.72, label="Scene preprocess")
    ax2.bar(
        xpos,
        frontend,
        bottom=preprocess,
        color=SEGMENT_COLORS["Front-end"],
        width=0.72,
        label="Front-end",
    )
    ax2.bar(
        xpos,
        backend,
        bottom=preprocess + frontend,
        color=SEGMENT_COLORS["Back-end"],
        width=0.72,
        label="Back-end",
    )
    ax2.set_xticks(xpos)
    ax2.set_xticklabels(
        ["Baseline", "No RS-MRQ", "No RV", "Strict No Mode", "Ours"],
        rotation=18,
        ha="center",
    )
    ax2.set_ylabel("Mean Time (s)")
    ax2.set_title("Timing Breakdown")
    ax2.legend(frameon=False, loc="upper left")
    fig.subplots_adjust(left=0.08, right=0.995, top=0.90, bottom=0.31, wspace=0.28)
    add_subfigure_labels_below(fig, axes, ["(a)", "(b)"], fixed_y=0.035)
    save_figure(fig, FIG_DIR / "main_ablation_figure")


def plot_adaptive_efficiency_figure(summary: pd.DataFrame) -> None:
    methods = ["No RS-MRQ", "Adaptive Two-Stage", "Ours (Full)"]
    fig, axes = plt.subplots(1, 2, figsize=(7.1, 2.85))
    label_text = {"No RS-MRQ": "No RS-MRQ", "Adaptive Two-Stage": "Adaptive", "Ours (Full)": "Ours"}
    label_offsets = {
        "Stanford": {
            "No RS-MRQ": (0.12, 0.0014, "left"),
            "Adaptive Two-Stage": (0.12, 0.0014, "left"),
            "Ours (Full)": (0.12, 0.0007, "left"),
        },
        "LMO": {
            "No RS-MRQ": (0.08, 0.0010, "left"),
            "Adaptive Two-Stage": (0.08, 0.0010, "left"),
            "Ours (Full)": (-0.08, 0.0010, "right"),
        },
    }
    for ax, dataset in zip(axes, ["Stanford", "LMO"], strict=True):
        sdf = summary_lookup(summary, dataset).loc[methods].reset_index()
        x = sdf["mean_registration_time"].to_numpy(dtype=float)
        ycol = "SR@ADD<=0.02" if dataset == "Stanford" else "primary_accuracy"
        ylabel = "SR@ADD<=0.02" if dataset == "Stanford" else "SR@ADD-S<=20mm"
        y = sdf[ycol].to_numpy(dtype=float)
        ax.set_xlim(*axis_limits(x, lower=0.0))
        ax.set_ylim(*axis_limits(y, lower=0.0, upper=1.0))
        for _, row in sdf.iterrows():
            method = row["method"]
            ax.scatter(
                row["mean_registration_time"],
                row[ycol],
                s=130 if method == "Adaptive Two-Stage" else 110,
                color=METHOD_COLORS[method],
                edgecolor="black" if method == "Adaptive Two-Stage" else "white",
                linewidth=1.0,
                zorder=3,
            )
            dx, dy, ha = label_offsets[dataset][method]
            ax.text(
                row["mean_registration_time"] + dx,
                row[ycol] + dy,
                label_text[method],
                ha=ha,
                va="bottom",
                fontsize=9,
                bbox=dict(facecolor="white", edgecolor="none", alpha=0.78, pad=0.1),
            )
        ax.set_title(dataset)
        ax.set_xlabel("Mean Registration Time (s)")
        ax.set_ylabel(ylabel)
    fig.subplots_adjust(left=0.09, right=0.995, top=0.88, bottom=0.29, wspace=0.24)
    add_subfigure_labels_below(fig, axes, ["(a)", "(b)"], fixed_y=0.04)
    save_figure(fig, FIG_DIR / "adaptive_efficiency_figure")


POINT_CACHE: dict[Path, np.ndarray] = {}


def load_points(path: Path | None) -> np.ndarray:
    if path is None:
        return np.zeros((0, 3), dtype=float)
    if path not in POINT_CACHE:
        pcd = o3d.io.read_point_cloud(str(path))
        POINT_CACHE[path] = np.asarray(pcd.points, dtype=np.float64)
    return POINT_CACHE[path]


def sample_points(points: np.ndarray, limit: int, seed: int = 0) -> np.ndarray:
    if points.shape[0] <= limit:
        return points
    rng = np.random.default_rng(seed)
    idx = rng.choice(points.shape[0], size=limit, replace=False)
    return points[idx]


def transform_points(points: np.ndarray, T: np.ndarray) -> np.ndarray:
    homog = np.concatenate([points, np.ones((points.shape[0], 1), dtype=np.float64)], axis=1)
    return (homog @ T.T)[:, :3]


def set_equal_3d(ax, points: np.ndarray) -> None:
    mins = points.min(axis=0)
    maxs = points.max(axis=0)
    center = (mins + maxs) / 2.0
    radius = max(float(np.max(maxs - mins)) / 2.0, 1e-6)
    ax.set_xlim(center[0] - radius, center[0] + radius)
    ax.set_ylim(center[1] - radius, center[1] + radius)
    ax.set_zlim(center[2] - radius, center[2] + radius)


def render_alignment_on_axis(ax, scene_path: Path | None, model_path: Path | None, T: np.ndarray, panel_title: str, metric_text: str) -> None:
    scene_pts = load_points(scene_path)
    model_pts = load_points(model_path)
    if scene_pts.size == 0 or model_pts.size == 0:
        ax.set_axis_off()
        ax.set_title(panel_title, fontsize=9)
        return
    scene_vis = sample_points(scene_pts, 3500, seed=0)
    model_vis = sample_points(transform_points(model_pts, T), 1600, seed=1)
    all_pts = np.concatenate([scene_vis, model_vis], axis=0)
    ax.scatter(scene_vis[:, 0], scene_vis[:, 1], scene_vis[:, 2], s=0.55, c="#A9B0B4", alpha=0.38)
    ax.scatter(model_vis[:, 0], model_vis[:, 1], model_vis[:, 2], s=0.9, c="#D94841", alpha=0.88)
    ax.view_init(elev=24, azim=45)
    set_equal_3d(ax, all_pts)
    ax.set_xticks([])
    ax.set_yticks([])
    ax.set_zticks([])
    title = metric_text if not panel_title else f"{panel_title}\n{metric_text}"
    ax.set_title(title, fontsize=8)


def select_stanford_qualitative_cases(master: pd.DataFrame) -> list[dict[str, Any]]:
    stanford = master[master["dataset"] == "Stanford"].copy()
    pivot = stanford.pivot_table(index="case_id", columns="method", values="success_add_002", aggfunc="first")
    add_pivot = stanford.pivot_table(index="case_id", columns="method", values="ADD", aggfunc="first")
    meta = stanford.drop_duplicates("case_id").set_index("case_id")

    chosen: list[dict[str, Any]] = []

    cat1 = pivot[(pivot["Ours (Full)"] == True) & (pivot["Same-backbone Baseline"] == False)]
    if not cat1.empty:
        best_case = (
            (add_pivot.loc[cat1.index, "Same-backbone Baseline"] - add_pivot.loc[cat1.index, "Ours (Full)"])
            .sort_values(ascending=False)
            .index[0]
        )
        chosen.append(
            {
                "case_id": best_case,
                "label": "Ours rescues weak baseline",
                "meta": meta.loc[best_case],
            }
        )

    cat2 = pivot[
        (pivot["Ours (Full)"] == True)
        & (pivot["No RS-MRQ"] == False)
        & (pivot["Adaptive Two-Stage"] == True)
    ]
    if not cat2.empty:
        best_case = (
            (add_pivot.loc[cat2.index, "No RS-MRQ"] - add_pivot.loc[cat2.index, "Adaptive Two-Stage"])
            .sort_values(ascending=False)
            .index[0]
        )
        chosen.append(
            {
                "case_id": best_case,
                "label": "Adaptive rescues No RS-MRQ",
                "meta": meta.loc[best_case],
            }
        )

    adaptive_up = stanford[
        (stanford["method"] == "Adaptive Two-Stage")
        & (stanford["adaptive_used_upgrade"] == True)
        & (stanford["success_add_002"] == True)
        & (~stanford["case_id"].isin([c["case_id"] for c in chosen]))
    ].sort_values("ADD")
    if not adaptive_up.empty:
        row = adaptive_up.iloc[0]
        chosen.append(
            {
                "case_id": row["case_id"],
                "label": "Adaptive upgraded hard case",
                "meta": row,
            }
        )
    extra = stanford[
        (stanford["method"] == "Ours (Full)")
        & (stanford["success_add_002"] == True)
        & (~stanford["case_id"].isin([c["case_id"] for c in chosen]))
    ].sort_values("ADD")
    if not extra.empty and len(chosen) < 4:
        row = extra.iloc[0]
        chosen.append({"case_id": row["case_id"], "label": "Representative successful case", "meta": row})
    return chosen


def plot_qualitative_rescue_failure(master: pd.DataFrame, render_lookup: dict[tuple[str, str, str], dict[str, Any]]) -> None:
    cases = select_stanford_qualitative_cases(master)[:3]
    methods = ["Same-backbone Baseline", "No RS-MRQ", "Adaptive Two-Stage", "Ours (Full)"]
    nrows = len(cases)
    header_text = {
        "Same-backbone Baseline": "Baseline",
        "No RS-MRQ": "No RS-MRQ",
        "Adaptive Two-Stage": "Adaptive",
        "Ours (Full)": "Ours",
    }
    fig = plt.figure(figsize=(10.4, 2.6 * max(nrows, 1)))
    gs = GridSpec(nrows, len(methods), figure=fig, wspace=0.02, hspace=0.20)
    for c, method in enumerate(methods):
        fig.text(0.17 + c * 0.205, 0.965, header_text[method], ha="center", va="top", fontsize=10, fontweight="bold")
    for r, case in enumerate(cases):
        meta = case["meta"]
        row_title = f"{case['label']}\n{meta['scene_name']} | v={meta['scene_variant']} | obj={meta['obj_id']}"
        y = 0.87 - r * (0.83 / max(nrows, 1))
        fig.text(0.015, y, row_title, va="center", ha="left", fontsize=8.7)
        for c, method in enumerate(methods):
            ax = fig.add_subplot(gs[r, c], projection="3d")
            panel = render_lookup.get(("Stanford", method, case["case_id"]))
            if panel is None:
                ax.set_axis_off()
                continue
            metric_text = f"ADD={panel['ADD']:.4f}" if np.isfinite(panel["ADD"]) else "ADD=nan"
            render_alignment_on_axis(ax, panel["scene_path"], panel["model_path"], panel["T_pred"], "", metric_text)
    fig.subplots_adjust(left=0.12, right=0.995, top=0.92, bottom=0.03)
    save_figure(fig, FIG_DIR / "qualitative_rescue_failure")


def select_lmo_mechanism_cases(master: pd.DataFrame, top_k: int = 3) -> list[str]:
    lmo = master[master["dataset"] == "LMO"]
    pivot_success = lmo.pivot_table(index="case_id", columns="method", values="primary_success", aggfunc="first")
    pivot_adds = lmo.pivot_table(index="case_id", columns="method", values="ADD_S", aggfunc="first")
    candidate = pivot_success[
        (pivot_success["Ours (Full)"] == True)
        & ((pivot_success["Same-backbone Baseline"] == False) | (pivot_success["No Robust Vote"] == False))
    ]
    if candidate.empty:
        return list(pivot_success.index[:top_k])
    score = (
        pivot_adds.loc[candidate.index, "Same-backbone Baseline"].fillna(0)
        + pivot_adds.loc[candidate.index, "No Robust Vote"].fillna(0)
        - 2 * pivot_adds.loc[candidate.index, "Ours (Full)"].fillna(0)
    )
    return score.sort_values(ascending=False).head(top_k).index.tolist()


def plot_lmo_matching_mechanism(master: pd.DataFrame, render_lookup: dict[tuple[str, str, str], dict[str, Any]]) -> None:
    cases = select_lmo_mechanism_cases(master, top_k=3)
    methods = ["Same-backbone Baseline", "No Robust Vote", "Ours (Full)", "GT"]
    lmo_meta = master[(master["dataset"] == "LMO") & (master["method"] == "Ours (Full)")].set_index("case_id")
    header_text = {
        "Same-backbone Baseline": "Baseline",
        "No Robust Vote": "No RV",
        "Ours (Full)": "Ours",
        "GT": "GT",
    }
    fig = plt.figure(figsize=(10.4, 2.7 * max(len(cases), 1)))
    gs = GridSpec(len(cases), len(methods), figure=fig, wspace=0.03, hspace=0.20)
    for c, method in enumerate(methods):
        fig.text(0.17 + c * 0.205, 0.965, header_text[method], ha="center", va="top", fontsize=10, fontweight="bold")
    for r, case_id in enumerate(cases):
        meta = lmo_meta.loc[case_id]
        row_title = f"frame={meta['frame_id']} | obj={meta['obj_id']}"
        y = 0.87 - r * (0.83 / max(len(cases), 1))
        fig.text(0.015, y, row_title, va="center", ha="left", fontsize=8.7)
        for c, method in enumerate(methods):
            ax = fig.add_subplot(gs[r, c], projection="3d")
            if method == "GT":
                panel = render_lookup.get(("LMO", "Ours (Full)", case_id))
                if panel is None:
                    ax.set_axis_off()
                    continue
                metric_text = "GT pose"
                render_alignment_on_axis(ax, panel["scene_path"], panel["model_path"], panel["T_gt"], "", metric_text)
                continue
            panel = render_lookup.get(("LMO", method, case_id))
            if panel is None:
                ax.set_axis_off()
                continue
            metric_text = f"ADD-S={panel['ADD_S']:.2f}" if np.isfinite(panel["ADD_S"]) else "ADD-S=nan"
            render_alignment_on_axis(ax, panel["scene_path"], panel["model_path"], panel["T_pred"], "", metric_text)
    fig.subplots_adjust(left=0.12, right=0.995, top=0.92, bottom=0.03)
    save_figure(fig, FIG_DIR / "lmo_matching_mechanism")


def plot_parameter_tradeoff(summary: pd.DataFrame) -> None:
    s = summary_lookup(summary, "Stanford")
    cap_methods = ["Ours (Full)", "Ours + RS-MRQ cap64", "Ours + RS-MRQ cap128", "Ours + RS-MRQ cap256"]
    rv_methods = ["Ours (Full)", "Ours + RV20", "Ours + RV10"]
    fig, axes = plt.subplots(1, 2, figsize=(7.1, 2.95))
    label_offsets = {
        "RS-MRQ Cap Trade-off": {
            "Ours (Full)": (6, -12),
            "Ours + RS-MRQ cap64": (8, 6),
            "Ours + RS-MRQ cap128": (8, 8),
            "Ours + RS-MRQ cap256": (8, 8),
        },
        "Robust Vote Light Trade-off": {
            "Ours (Full)": (8, -12),
            "Ours + RV20": (8, 8),
            "Ours + RV10": (8, 8),
        },
    }
    for ax, methods, title in zip(
        axes,
        [cap_methods, rv_methods],
        ["RS-MRQ Cap Trade-off", "Robust Vote Light Trade-off"],
        strict=True,
    ):
        df = s.loc[methods].reset_index()
        ax.plot(df["mean_registration_time"], df["mean_ADD"], color="#666666", lw=1.2, zorder=1)
        for _, row in df.iterrows():
            method = row["method"]
            ax.scatter(
                row["mean_registration_time"],
                row["mean_ADD"],
                s=120 if method == "Ours (Full)" else 100,
                color=METHOD_COLORS[method],
                edgecolor="black" if method == "Ours (Full)" else "white",
                linewidth=1.0,
                zorder=3,
            )
            short = (
                "original"
                if method == "Ours (Full)"
                else method.replace("Ours + RS-MRQ ", "").replace("Ours + ", "")
            )
            offset = label_offsets[title][method]
            ax.annotate(
                short,
                xy=(row["mean_registration_time"], row["mean_ADD"]),
                xytext=offset,
                textcoords="offset points",
                fontsize=8,
                bbox=dict(facecolor="white", edgecolor="none", alpha=0.78, pad=0.1),
            )
        ax.set_title(title)
        ax.set_xlabel("Mean Registration Time (s)")
        ax.set_ylabel("Mean ADD")
    fig.subplots_adjust(left=0.09, right=0.995, top=0.90, bottom=0.31, wspace=0.34)
    add_subfigure_labels_below(fig, axes, ["(a)", "(b)"], fixed_y=0.035)
    save_figure(fig, FIG_DIR / "parameter_tradeoff_overview")


def plot_stanford_efficiency_parameter_summary(summary: pd.DataFrame) -> None:
    s = summary_lookup(summary, "Stanford")
    eff_methods = ["No RS-MRQ", "Ours (Full)", "Adaptive Two-Stage", "Ours + RV20", "Ours + RV10", "Ours + RS-MRQ cap256"]
    param_methods = ["Ours (Full)", "Ours + RS-MRQ cap256", "Ours + RS-MRQ cap128", "Ours + RS-MRQ cap64", "Ours + RV20", "Ours + RV10", "Adaptive Two-Stage"]
    label_map = {
        "No RS-MRQ": "No RS-MRQ",
        "Ours (Full)": "Original",
        "Adaptive Two-Stage": "Adaptive",
        "Ours + RV20": "RV20",
        "Ours + RV10": "RV10",
        "Ours + RS-MRQ cap256": "cap256",
        "Ours + RS-MRQ cap128": "cap128",
        "Ours + RS-MRQ cap64": "cap64",
    }
    eff_label_layout = {
        "No RS-MRQ": (10, 12, "left", "bottom"),
        "Adaptive Two-Stage": (-10, 8, "right", "bottom"),
        "Ours + RV10": (0, 10, "center", "bottom"),
        "Ours + RV20": (10, 8, "left", "bottom"),
        "Ours (Full)": (-8, 10, "right", "bottom"),
        "Ours + RS-MRQ cap256": (8, 8, "left", "bottom"),
    }
    param_label_layout = {
        "Adaptive Two-Stage": (-10, 8, "right", "bottom"),
        "Ours + RV10": (0, 10, "center", "bottom"),
        "Ours + RV20": (10, 8, "left", "bottom"),
        "Ours (Full)": (-8, 10, "right", "bottom"),
        "Ours + RS-MRQ cap256": (0, 10, "center", "bottom"),
        "Ours + RS-MRQ cap128": (8, 8, "left", "bottom"),
        "Ours + RS-MRQ cap64": (-8, -10, "right", "top"),
    }
    selected_methods = {"Ours (Full)", "Adaptive Two-Stage", "Ours + RV20", "Ours + RS-MRQ cap256"}
    fig, axes = plt.subplots(1, 2, figsize=(7.3, 3.05))

    eff_df = s.loc[eff_methods].reset_index()
    x = eff_df["mean_registration_time"].to_numpy(dtype=float)
    y = eff_df["SR@ADD<=0.02"].to_numpy(dtype=float)
    ax = axes[0]
    ax.set_xlim(*axis_limits(x, lower=0.0))
    ax.set_ylim(*axis_limits(y, lower=0.0, upper=1.0))
    ax.plot(x, y, color="#B7B7B7", lw=1.1, zorder=1)
    for _, row in eff_df.iterrows():
        method = row["method"]
        ax.scatter(
            row["mean_registration_time"],
            row["SR@ADD<=0.02"],
            s=125 if method in selected_methods else 105,
            color=METHOD_COLORS[method],
            edgecolor="black" if method in selected_methods else "white",
            linewidth=1.0,
            zorder=3,
        )
        dx, dy, ha, va = eff_label_layout[method]
        ax.annotate(
            label_map[method],
            xy=(row["mean_registration_time"], row["SR@ADD<=0.02"]),
            xytext=(dx, dy),
            textcoords="offset points",
            ha=ha,
            va=va,
            fontsize=8.5,
            annotation_clip=False,
            arrowprops=dict(arrowstyle="-", color="#8C8C8C", lw=0.55, shrinkA=3, shrinkB=4),
            bbox=dict(facecolor="white", edgecolor="none", alpha=0.8, pad=0.12),
        )
    ax.set_title("Efficiency extensions")
    ax.set_xlabel("Mean registration time (s)")
    ax.set_ylabel("SR@ADD<=0.02")

    ax = axes[1]
    cap_df = s.loc[["Ours (Full)", "Ours + RS-MRQ cap256", "Ours + RS-MRQ cap128", "Ours + RS-MRQ cap64"]].reset_index()
    rv_df = s.loc[["Ours (Full)", "Ours + RV20", "Ours + RV10", "Adaptive Two-Stage"]].reset_index()
    ax.plot(cap_df["mean_registration_time"], cap_df["SR@ADD<=0.02"], color="#8C8C8C", lw=1.0, zorder=1)
    ax.plot(rv_df["mean_registration_time"], rv_df["SR@ADD<=0.02"], color="#8C8C8C", lw=1.0, ls="--", zorder=1)
    param_df = s.loc[param_methods].reset_index()
    x = param_df["mean_registration_time"].to_numpy(dtype=float)
    y = param_df["SR@ADD<=0.02"].to_numpy(dtype=float)
    ax.set_xlim(*axis_limits(x, lower=0.0))
    ax.set_ylim(*axis_limits(y, lower=0.0, upper=1.0))
    for _, row in param_df.iterrows():
        method = row["method"]
        marker = "s" if "cap" in method else "^" if ("RV" in method or method == "Adaptive Two-Stage") else "o"
        ax.scatter(
            row["mean_registration_time"],
            row["SR@ADD<=0.02"],
            s=125 if method in selected_methods else 102,
            color=METHOD_COLORS[method],
            marker=marker,
            edgecolor="black" if method in selected_methods else "white",
            linewidth=1.0,
            zorder=3,
        )
        dx, dy, ha, va = param_label_layout[method]
        ax.annotate(
            label_map[method],
            xy=(row["mean_registration_time"], row["SR@ADD<=0.02"]),
            xytext=(dx, dy),
            textcoords="offset points",
            ha=ha,
            va=va,
            fontsize=8.5,
            annotation_clip=False,
            arrowprops=dict(arrowstyle="-", color="#8C8C8C", lw=0.55, shrinkA=3, shrinkB=4),
            bbox=dict(facecolor="white", edgecolor="none", alpha=0.8, pad=0.12),
        )
    ax.set_title("Parameter selection")
    ax.set_xlabel("Mean registration time (s)")
    ax.set_ylabel("SR@ADD<=0.02")

    fig.subplots_adjust(left=0.09, right=0.995, top=0.89, bottom=0.29, wspace=0.28)
    add_subfigure_labels_below(fig, axes, ["(a)", "(b)"], fixed_y=0.045)
    save_figure(fig, FIG_DIR / "stanford_efficiency_parameter_summary")


def build_appendix_tables(master: pd.DataFrame, summary: pd.DataFrame) -> None:
    stanford_summary = summary[summary["dataset"] == "Stanford"].sort_values("method_order")
    lmo_summary = summary[summary["dataset"] == "LMO"].sort_values("method_order")
    stanford_ablation = stanford_summary[stanford_summary["method"].isin(CORE_METHODS)]
    lmo_ablation = lmo_summary[lmo_summary["method"].isin(CORE_METHODS)]
    adaptive_detail = {
        "Adaptive_Summary": pd.read_csv(STANFORD_ANALYSIS_DIR / "adaptive_efficiency_table.csv"),
        "Adaptive_TriggerBreakdown": pd.read_csv(STANFORD_ANALYSIS_DIR / "adaptive_upgrade_breakdown.csv"),
        "Adaptive_ByScene": pd.read_csv(STANFORD_ANALYSIS_DIR / "adaptive_upgrade_by_scene.csv"),
        "Adaptive_ByObject": pd.read_csv(STANFORD_ANALYSIS_DIR / "adaptive_upgrade_by_object.csv"),
        "Adaptive_Rescue": pd.read_csv(STANFORD_ANALYSIS_DIR / "adaptive_rescue_summary.csv"),
    }
    rsmrq_full = pd.read_csv(STANFORD_ANALYSIS_DIR / "rsmrq_budget_sweep_enriched.csv")
    rv_full = pd.read_csv(STANFORD_ANALYSIS_DIR / "robust_vote_light_enriched.csv")
    hard_case_master = pd.read_csv(STANFORD_ANALYSIS_DIR / "hard_case_master.csv")

    lmo_time_tail = build_time_tail_table(master[master["dataset"] == "LMO"])
    stanford_time_tail = pd.read_csv(STANFORD_ANALYSIS_DIR / "time_tail_statistics.csv")
    by_scene = pd.read_csv(STANFORD_ANALYSIS_DIR / "by_scene_summary.csv")
    by_object = pd.read_csv(STANFORD_ANALYSIS_DIR / "by_object_summary.csv")
    lmo_by_object = build_lmo_by_object_summary(master)
    lmo_by_scene = build_lmo_by_scene_summary(master)
    df_to_workbook(APP_TABLE_DIR / "stanford_detailed_comparison.xlsx", {"Stanford_Comparison": stanford_summary})
    df_to_workbook(APP_TABLE_DIR / "lmo_detailed_comparison.xlsx", {"LMO_Comparison": lmo_summary})
    df_to_workbook(APP_TABLE_DIR / "stanford_detailed_ablation.xlsx", {"Stanford_Ablation": stanford_ablation})
    df_to_workbook(APP_TABLE_DIR / "lmo_detailed_ablation.xlsx", {"LMO_Ablation": lmo_ablation})
    df_to_workbook(APP_TABLE_DIR / "adaptive_detailed.xlsx", adaptive_detail)
    df_to_workbook(APP_TABLE_DIR / "rsmrq_budget_sweep_full.xlsx", {"RSMRQ_Sweep": rsmrq_full})
    df_to_workbook(APP_TABLE_DIR / "robust_vote_light_sweep_full.xlsx", {"RV_Sweep": rv_full})
    df_to_workbook(APP_TABLE_DIR / "hard_case_master.xlsx", {"HardCases": hard_case_master})
    df_to_workbook(
        APP_TABLE_DIR / "time_tail_statistics.xlsx",
        {"Stanford_Tail": stanford_time_tail, "LMO_Tail": lmo_time_tail},
    )
    df_to_workbook(
        APP_TABLE_DIR / "by_scene_by_object_summary.xlsx",
        {
            "Stanford_ByScene": by_scene,
            "Stanford_ByObject": by_object,
            "LMO_ByScene": lmo_by_scene,
            "LMO_ByObject": lmo_by_object,
        },
    )


def build_time_tail_table(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for method, gdf in df.groupby("method", sort=False):
        vals = pd.to_numeric(gdf["registration_time"], errors="coerce").dropna().to_numpy(dtype=float)
        if vals.size == 0:
            continue
        rows.append(
            {
                "method": method,
                "mean": float(np.mean(vals)),
                "median": float(np.median(vals)),
                "p90": float(np.percentile(vals, 90)),
                "p95": float(np.percentile(vals, 95)),
                "p99": float(np.percentile(vals, 99)),
                "max": float(np.max(vals)),
                "n_cases_gt_10s": int(np.sum(vals > 10.0)),
                "n_cases_gt_30s": int(np.sum(vals > 30.0)),
                "n_cases_gt_60s": int(np.sum(vals > 60.0)),
            }
        )
    return pd.DataFrame(rows).sort_values(
        by="method", key=lambda s: s.map(METHOD_ORDER_INDEX).fillna(999)
    )


def build_lmo_by_object_summary(master: pd.DataFrame) -> pd.DataFrame:
    lmo = master[master["dataset"] == "LMO"]
    rows = []
    for (method, obj_id), gdf in lmo.groupby(["method", "obj_id"], sort=False):
        rows.append(
            {
                "method": method,
                "obj_id": obj_id,
                "Primary Accuracy Metric": float(pd.Series(gdf["primary_success"]).mean()),
                "Mean ADD-S": nanmean(gdf["ADD_S"]),
                "Mean registration_time": nanmean(gdf["registration_time"]),
                "Mean frontend_time": nanmean(gdf["ppf_frontend_time"]),
            }
        )
    return pd.DataFrame(rows).sort_values(
        ["method", "obj_id"], key=lambda s: s.map(METHOD_ORDER_INDEX).fillna(999) if s.name == "method" else s
    )


def build_lmo_by_scene_summary(master: pd.DataFrame) -> pd.DataFrame:
    lmo = master[master["dataset"] == "LMO"]
    rows = []
    for (method, scene_name), gdf in lmo.groupby(["method", "scene_name"], sort=False):
        rows.append(
            {
                "method": method,
                "scene_name": scene_name,
                "Primary Accuracy Metric": float(pd.Series(gdf["primary_success"]).mean()),
                "Mean ADD-S": nanmean(gdf["ADD_S"]),
                "Mean registration_time": nanmean(gdf["registration_time"]),
                "Mean frontend_time": nanmean(gdf["ppf_frontend_time"]),
            }
        )
    return pd.DataFrame(rows).sort_values(
        ["method", "scene_name"], key=lambda s: s.map(METHOD_ORDER_INDEX).fillna(999) if s.name == "method" else s
    )


def plot_success_curves(master: pd.DataFrame, dataset: str, out_name: str) -> None:
    data = master[master["dataset"] == dataset]
    if dataset == "Stanford":
        metric = "ADD"
        thresholds = np.linspace(0.0, 0.05, 80)
        ylabel = "Success Rate"
    else:
        metric = "ADD_S"
        thresholds = np.linspace(0.0, 60.0, 80)
        ylabel = "Success Rate"
    fig, axes = plt.subplots(1, 2, figsize=(10.0, 4.0))
    internal = CORE_METHODS + ["Adaptive Two-Stage"]
    external = ["Same-backbone Baseline", "No RS-MRQ", "Ours (Full)"] + EXTERNAL_METHODS
    for ax, methods, title in zip(axes, [internal, external], ["Internal Variants", "External Baselines"], strict=True):
        for method in methods:
            sub = data[data["method"] == method]
            values = pd.to_numeric(sub[metric], errors="coerce").dropna().to_numpy(dtype=float)
            if values.size == 0:
                continue
            curve = [(values <= t).mean() for t in thresholds]
            ax.plot(thresholds, curve, lw=2.0 if method == "Ours (Full)" else 1.5, color=METHOD_COLORS[method], label=method)
        ax.set_title(title)
        ax.set_xlabel("ADD threshold" if dataset == "Stanford" else "ADD-S threshold (mm)")
        ax.set_ylabel(ylabel)
        ax.legend(frameon=False, fontsize=8)
    fig.suptitle(f"{dataset} Success Curves", y=1.02, fontsize=14)
    fig.tight_layout()
    save_figure(fig, APP_FIG_DIR / out_name)


def plot_time_distribution(master: pd.DataFrame, dataset: str, out_name: str) -> None:
    data = master[master["dataset"] == dataset]
    methods = ["Same-backbone Baseline", "No RS-MRQ", "No Robust Vote", "Ours (Full)", "Adaptive Two-Stage", "Drost", "Going Further"]
    fig, axes = plt.subplots(1, 2, figsize=(11.2, 4.6))
    samples = [pd.to_numeric(data[data["method"] == m]["registration_time"], errors="coerce").dropna().to_numpy(dtype=float) for m in methods]
    box = axes[0].boxplot(samples, tick_labels=methods, patch_artist=True)
    for patch, method in zip(box["boxes"], methods):
        patch.set_facecolor(METHOD_COLORS[method])
    axes[0].set_title("Boxplot")
    axes[0].set_ylabel("Registration Time (s)")
    axes[0].tick_params(axis="x", rotation=25)
    parts = axes[1].violinplot(samples, showmeans=True, showextrema=False)
    for body, method in zip(parts["bodies"], methods):
        body.set_facecolor(METHOD_COLORS[method])
        body.set_alpha(0.55)
    axes[1].set_xticks(np.arange(1, len(methods) + 1))
    axes[1].set_xticklabels(methods, rotation=25, ha="right")
    axes[1].set_title("Violin")
    axes[1].set_ylabel("Registration Time (s)")
    fig.suptitle(f"{dataset} Time Distribution", y=1.02, fontsize=14)
    fig.tight_layout()
    save_figure(fig, APP_FIG_DIR / out_name)


def plot_adaptive_trigger_breakdown() -> None:
    reason_df = pd.read_csv(STANFORD_ANALYSIS_DIR / "adaptive_upgrade_breakdown.csv")
    scene_df = pd.read_csv(STANFORD_ANALYSIS_DIR / "adaptive_upgrade_by_scene.csv").head(10)
    object_df = pd.read_csv(STANFORD_ANALYSIS_DIR / "adaptive_upgrade_by_object.csv").head(10)
    fig, axes = plt.subplots(1, 3, figsize=(12.5, 4.0))
    axes[0].bar(reason_df["trigger_reason"], reason_df["case_ratio"], color="#4E79A7")
    axes[0].set_title("Trigger Reason Share")
    axes[0].tick_params(axis="x", rotation=25)
    axes[0].set_ylabel("Case Ratio")
    axes[1].bar(scene_df["scene_name"], scene_df["upgrade_ratio"], color="#59A14F")
    axes[1].set_title("Upgrade Ratio by Scene")
    axes[1].tick_params(axis="x", rotation=25)
    axes[2].bar(object_df["obj_id"].astype(str), object_df["upgrade_ratio"], color="#F28E2B")
    axes[2].set_title("Upgrade Ratio by Object")
    axes[2].tick_params(axis="x", rotation=25)
    fig.suptitle("Adaptive Trigger Breakdown", y=1.02, fontsize=14)
    fig.tight_layout()
    save_figure(fig, APP_FIG_DIR / "adaptive_trigger_breakdown")


def plot_by_scene_analysis(master: pd.DataFrame) -> None:
    stanford = master[master["dataset"] == "Stanford"]
    methods = ["No RS-MRQ", "Adaptive Two-Stage", "Ours (Full)"]
    grouped = (
        stanford[stanford["method"].isin(methods)]
        .groupby(["scene_name", "method"], sort=False)
        .agg(success=("success_add_002", "mean"), reg=("registration_time", "mean"))
        .reset_index()
    )
    scenes = sorted(grouped["scene_name"].unique(), key=lambda x: int(re.findall(r"\d+", str(x))[0]) if re.findall(r"\d+", str(x)) else 0)
    fig, axes = plt.subplots(1, 2, figsize=(11.0, 4.2))
    for method in methods:
        sub = grouped[grouped["method"] == method].set_index("scene_name").reindex(scenes)
        axes[0].plot(scenes, sub["success"], marker="o", lw=1.8, color=METHOD_COLORS[method], label=method)
        axes[1].plot(scenes, sub["reg"], marker="o", lw=1.8, color=METHOD_COLORS[method], label=method)
    axes[0].set_title("Stanford by Scene: Success")
    axes[0].set_ylabel("SR@ADD<=0.02")
    axes[0].tick_params(axis="x", rotation=25)
    axes[1].set_title("Stanford by Scene: Runtime")
    axes[1].set_ylabel("Mean Registration Time (s)")
    axes[1].tick_params(axis="x", rotation=25)
    axes[0].legend(frameon=False, loc="lower right")
    fig.tight_layout()
    save_figure(fig, APP_FIG_DIR / "by_scene_analysis")


def plot_by_object_analysis(master: pd.DataFrame) -> None:
    s_methods = ["No RS-MRQ", "Adaptive Two-Stage", "Ours (Full)"]
    stanford = (
        master[(master["dataset"] == "Stanford") & (master["method"].isin(s_methods))]
        .groupby(["obj_id", "method"], sort=False)
        .agg(success=("success_add_002", "mean"))
        .reset_index()
    )
    lmo = (
        master[(master["dataset"] == "LMO") & (master["method"].isin(s_methods))]
        .groupby(["obj_id", "method"], sort=False)
        .agg(success=("primary_success", "mean"))
        .reset_index()
    )
    fig, axes = plt.subplots(1, 2, figsize=(11.0, 4.2))
    for method in s_methods:
        ssub = stanford[stanford["method"] == method]
        axes[0].plot(ssub["obj_id"].astype(str), ssub["success"], marker="o", lw=1.6, color=METHOD_COLORS[method], label=method)
        lsub = lmo[lmo["method"] == method]
        axes[1].plot(lsub["obj_id"].astype(str), lsub["success"], marker="o", lw=1.6, color=METHOD_COLORS[method], label=method)
    axes[0].set_title("Stanford by Object")
    axes[0].set_ylabel("Success Rate")
    axes[1].set_title("LMO by Object")
    axes[1].set_ylabel("SR@ADD-S<=20mm")
    axes[0].legend(frameon=False, loc="lower right")
    fig.tight_layout()
    save_figure(fig, APP_FIG_DIR / "by_object_analysis")


def plot_lmo_matching_gallery(master: pd.DataFrame, render_lookup: dict[tuple[str, str, str], dict[str, Any]]) -> None:
    cases = select_lmo_mechanism_cases(master, top_k=4)
    methods = ["Same-backbone Baseline", "No RS-MRQ", "No Robust Vote", "Adaptive Two-Stage", "Ours (Full)"]
    fig = plt.figure(figsize=(12.5, 3.0 * len(cases)))
    gs = GridSpec(len(cases), len(methods), figure=fig, wspace=0.03, hspace=0.20)
    meta_lookup = master[(master["dataset"] == "LMO") & (master["method"] == "Ours (Full)")].set_index("case_id")
    for r, case_id in enumerate(cases):
        meta = meta_lookup.loc[case_id]
        fig.text(0.01, 1 - (r + 0.5) / len(cases), f"frame={meta['frame_id']} | obj={meta['obj_id']}", ha="left", va="center", fontsize=9)
        for c, method in enumerate(methods):
            ax = fig.add_subplot(gs[r, c], projection="3d")
            panel = render_lookup.get(("LMO", method, case_id))
            if panel is None:
                ax.set_axis_off()
                continue
            metric_text = f"ADD-S={panel['ADD_S']:.2f}" if np.isfinite(panel["ADD_S"]) else "ADD-S=nan"
            render_alignment_on_axis(ax, panel["scene_path"], panel["model_path"], panel["T_pred"], method, metric_text)
    fig.suptitle("LMO Matching Gallery", fontsize=14, y=0.99)
    fig.subplots_adjust(left=0.11, right=0.995, top=0.94, bottom=0.03)
    save_figure(fig, APP_FIG_DIR / "lmo_matching_gallery")


def plot_rescue_failure_gallery(master: pd.DataFrame, render_lookup: dict[tuple[str, str, str], dict[str, Any]]) -> None:
    cases = select_stanford_qualitative_cases(master)[:4]
    methods = ["Same-backbone Baseline", "No RS-MRQ", "Adaptive Two-Stage", "Ours (Full)"]
    fig = plt.figure(figsize=(10.8, 3.0 * len(cases)))
    gs = GridSpec(len(cases), len(methods), figure=fig, wspace=0.03, hspace=0.20)
    for r, case in enumerate(cases):
        meta = case["meta"]
        fig.text(
            0.01,
            1 - (r + 0.5) / len(cases),
            f"{case['label']}\n{meta['scene_name']} | var={meta['scene_variant']} | obj={meta['obj_id']}",
            ha="left",
            va="center",
            fontsize=9,
        )
        for c, method in enumerate(methods):
            ax = fig.add_subplot(gs[r, c], projection="3d")
            panel = render_lookup.get(("Stanford", method, case["case_id"]))
            if panel is None:
                ax.set_axis_off()
                continue
            metric_text = f"ADD={panel['ADD']:.4f}" if np.isfinite(panel["ADD"]) else "ADD=nan"
            render_alignment_on_axis(ax, panel["scene_path"], panel["model_path"], panel["T_pred"], method, metric_text)
    fig.suptitle("Rescue / Failure Gallery", fontsize=14, y=0.99)
    fig.subplots_adjust(left=0.11, right=0.995, top=0.94, bottom=0.03)
    save_figure(fig, APP_FIG_DIR / "rescue_failure_gallery")


def plot_frontend_cost_driver_plots(master: pd.DataFrame) -> None:
    df = master[master["dataset"] == "Stanford"].copy()
    metrics = [
        ("candidate_inflation_mean", "Candidate Inflation"),
        ("num_input_candidates", "Num Input Candidates"),
        ("n_votes", "n_votes"),
    ]
    fig, axes = plt.subplots(1, 3, figsize=(12.5, 3.8))
    methods = ["Same-backbone Baseline", "No RS-MRQ", "No Robust Vote", "Ours (Full)", "Adaptive Two-Stage"]
    for ax, (column, title) in zip(axes, metrics, strict=True):
        for method in methods:
            sub = df[df["method"] == method]
            ax.scatter(
                sub[column],
                sub["ppf_frontend_time"],
                s=18,
                alpha=0.6,
                color=METHOD_COLORS[method],
                label=method if ax is axes[0] else None,
            )
        valid = df[[column, "ppf_frontend_time"]].dropna()
        if len(valid) > 1:
            coeffs = np.polyfit(valid[column], valid["ppf_frontend_time"], 1)
            xs = np.linspace(valid[column].min(), valid[column].max(), 100)
            ax.plot(xs, coeffs[0] * xs + coeffs[1], color="#333333", lw=1.0)
        ax.set_title(title)
        ax.set_xlabel(title)
        ax.set_ylabel("Front-end Time (s)")
    axes[0].legend(frameon=False, fontsize=8)
    fig.suptitle("Front-end Cost Driver Plots", y=1.02, fontsize=14)
    fig.tight_layout()
    save_figure(fig, APP_FIG_DIR / "frontend_cost_driver_plots")


def collect_main_and_appendix_files() -> list[tuple[str, str]]:
    manifest_items = []
    for folder, kind in [
        (TABLE_DIR, "正文表"),
        (FIG_DIR, "正文图"),
        (APP_TABLE_DIR, "附录表"),
        (APP_FIG_DIR, "附录图"),
        (ANALYSIS_DIR, "分析"),
    ]:
        if not folder.exists():
            continue
        for path in sorted(folder.rglob("*")):
            if path.is_file():
                manifest_items.append((str(path.relative_to(PAPER_DIR)), kind))
    return manifest_items


def build_deliverables_manifest() -> None:
    rows = []
    for rel_path, kind in collect_main_and_appendix_files():
        full_path = PAPER_DIR / rel_path
        generated = full_path.exists()
        quality_pass = generated
        notes = ""
        if full_path.suffix.lower() == ".png":
            try:
                with Image.open(full_path) as img:
                    dpi = img.info.get("dpi", (PNG_DPI, PNG_DPI))
                    quality_pass = quality_pass and dpi[0] >= 300
                    notes = f"dpi={dpi[0]:.0f}"
            except Exception as exc:
                quality_pass = False
                notes = f"dpi_check_failed: {exc}"
        elif full_path.suffix.lower() == ".svg":
            notes = "vector exported"
        rows.append(
            {
                "file_name": rel_path,
                "type": kind,
                "generated_successfully": generated,
                "quality_check_pass": quality_pass,
                "notes": notes,
            }
        )
    df_to_workbook(ANALYSIS_DIR / "deliverables_manifest.xlsx", {"Manifest": pd.DataFrame(rows)})


def main() -> None:
    ensure_dirs()
    apply_plot_style()
    master, summary, render_lookup = build_master_tables()
    build_master_outputs(master, summary)
    build_main_results_table(summary)
    build_overall_ablation_table(summary)
    build_efficiency_extensions_table(summary)
    build_parameter_selection_table(summary)
    plot_main_ablation_figure(summary)
    plot_adaptive_efficiency_figure(summary)
    plot_qualitative_rescue_failure(master, render_lookup)
    plot_lmo_matching_mechanism(master, render_lookup)
    plot_parameter_tradeoff(summary)
    plot_stanford_efficiency_parameter_summary(summary)
    build_appendix_tables(master, summary)
    plot_success_curves(master, "Stanford", "stanford_success_curves")
    plot_success_curves(master, "LMO", "lmo_success_curves")
    plot_time_distribution(master, "Stanford", "stanford_time_distribution")
    plot_time_distribution(master, "LMO", "lmo_time_distribution")
    plot_adaptive_trigger_breakdown()
    plot_by_scene_analysis(master)
    plot_by_object_analysis(master)
    plot_lmo_matching_gallery(master, render_lookup)
    plot_rescue_failure_gallery(master, render_lookup)
    plot_frontend_cost_driver_plots(master)
    build_deliverables_manifest()


if __name__ == "__main__":
    main()
