import csv
import json
from pathlib import Path
from statistics import mean
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[1]
RESULTS_DIR = REPO_ROOT / "experiments" / "results" / "Stanford"
OUT_DIR = REPO_ROOT / "experiments" / "tables" / "stanford" / "analysis_plus"
XLSX_OUT = OUT_DIR / "stanford_all_methods_comparison.xlsx"
CSV_OUT = OUT_DIR / "stanford_all_methods_comparison.csv"


def batch_json_path(filename: str) -> Path:
    preferred = RESULTS_DIR / filename
    if preferred.exists():
        return preferred
    return RESULTS_DIR.parent / filename


METHOD_SPECS = [
    {
        "label": "Same-backbone Baseline",
        "group": "Stanford Internal",
        "batch_json": batch_json_path("stanford_same_backbone_baseline_batch.json"),
    },
    {
        "label": "No RS-MRQ",
        "group": "Stanford Internal",
        "batch_json": batch_json_path("stanford_no_rsmrq_batch.json"),
    },
    {
        "label": "No Robust Vote",
        "group": "Stanford Internal",
        "batch_json": batch_json_path("stanford_no_robust_vote_batch.json"),
    },
    {
        "label": "No Backend",
        "group": "Stanford Internal",
        "batch_json": batch_json_path("stanford_no_backend_batch.json"),
    },
    {
        "label": "Ours (Full)",
        "group": "Stanford Internal",
        "batch_json": batch_json_path("stanford_ours_full_batch.json"),
    },
    {
        "label": "Adaptive Two-Stage",
        "group": "Stanford Internal",
        "batch_json": batch_json_path("stanford_adaptive_two_stage_batch.json"),
    },
    {
        "label": "Ours + RV20",
        "group": "Stanford Internal",
        "batch_json": batch_json_path("stanford_ours_rv20_batch.json"),
    },
    {
        "label": "Ours + RV10",
        "group": "Stanford Internal",
        "batch_json": batch_json_path("stanford_ours_rv10_batch.json"),
    },
    {
        "label": "Ours + RS-MRQ cap64",
        "group": "Stanford Internal",
        "batch_json": batch_json_path("stanford_ours_cap64_batch.json"),
    },
    {
        "label": "Ours + RS-MRQ cap128",
        "group": "Stanford Internal",
        "batch_json": batch_json_path("stanford_ours_cap128_batch.json"),
    },
    {
        "label": "Ours + RS-MRQ cap256",
        "group": "Stanford Internal",
        "batch_json": batch_json_path("stanford_ours_cap256_batch.json"),
    },
    {
        "label": "Drost original PPF",
        "group": "External PPF",
        "batch_json": batch_json_path("stanford_drost_original_batch.json"),
    },
    {
        "label": "Birdal Revisited",
        "group": "External PPF",
        "batch_json": batch_json_path("stanford_birdal_revisited_batch.json"),
    },
    {
        "label": "Going Further with PPF",
        "group": "External PPF",
        "batch_json": batch_json_path("stanford_going_further_ppf_batch.json"),
    },
    {
        "label": "Edge-enhanced PPF",
        "group": "External PPF",
        "batch_json": batch_json_path("stanford_edge_enhanced_ppf_batch.json"),
    },
]


def _safe_mean(values: Iterable[Optional[float]]) -> Optional[float]:
    vals = [float(v) for v in values if v is not None]
    if not vals:
        return None
    return mean(vals)


def _extract_value(result: Dict[str, Any], *names: str) -> Optional[float]:
    for container_name in ("metrics", "stats", "debug"):
        container = result.get(container_name, {}) or {}
        for name in names:
            if name in container and container[name] is not None:
                return float(container[name])
    return None


def _success_rate(result: Dict[str, Any], threshold: float = 0.02) -> Optional[float]:
    add = _extract_value(result, "ADD")
    if add is None:
        return None
    return 1.0 if add <= threshold else 0.0


def _aggregate_batch(spec: Dict[str, Any]) -> Dict[str, Any]:
    batch_json = spec["batch_json"]
    payload = json.loads(batch_json.read_text(encoding="utf-8"))
    results = payload.get("results", []) or []
    if not results:
        raise ValueError(f"No results found in {batch_json}")

    row = {
        "Method Group": spec["group"],
        "Method": spec["label"],
        "n_cases": len(results),
        "Mean ADD": _safe_mean(_extract_value(r, "ADD") for r in results),
        "Mean ADD-S": _safe_mean(_extract_value(r, "ADD-S", "ADD_S") for r in results),
        "SR@ADD<=0.02": _safe_mean(_success_rate(r) for r in results),
        "Mean Rotation Error (deg)": _safe_mean(_extract_value(r, "rotation_error_deg") for r in results),
        "Mean Translation Error": _safe_mean(_extract_value(r, "translation_error") for r in results),
        "Mean Registration Time (s)": _safe_mean(_extract_value(r, "registration_time") for r in results),
        "Mean Scene Preprocess Time (s)": _safe_mean(_extract_value(r, "scene_preprocess_time") for r in results),
        "Mean Front-end Time (s)": _safe_mean(_extract_value(r, "ppf_frontend_time") for r in results),
        "Mean Back-end Time (s)": _safe_mean(_extract_value(r, "backend_time") for r in results),
        "Candidate Inflation Mean": _safe_mean(_extract_value(r, "candidate_inflation_mean") for r in results),
        "Source Batch JSON": str(batch_json),
    }
    return row


def _write_csv(rows: List[Dict[str, Any]], columns: List[str]) -> None:
    with CSV_OUT.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def _auto_width(ws) -> None:
    for idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 12), 36)


def _write_workbook(rows: List[Dict[str, Any]], columns: List[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "All Methods"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    ours_fill = PatternFill(fill_type="solid", fgColor="FDE9D9")
    external_fill = PatternFill(fill_type="solid", fgColor="EAF4E3")

    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = Font(name="Times New Roman", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = header_fill

    numeric_formats = {
        "Mean ADD": "0.000000",
        "Mean ADD-S": "0.000000",
        "SR@ADD<=0.02": "0.000000",
        "Mean Rotation Error (deg)": "0.000000",
        "Mean Translation Error": "0.000000",
        "Mean Registration Time (s)": "0.000000",
        "Mean Scene Preprocess Time (s)": "0.000000",
        "Mean Front-end Time (s)": "0.000000",
        "Mean Back-end Time (s)": "0.000000",
        "Candidate Inflation Mean": "0.000000",
    }

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(columns, start=1):
            value = row.get(col_name)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name="Times New Roman")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_name in numeric_formats and isinstance(value, (int, float)):
                cell.number_format = numeric_formats[col_name]

        if row["Method"] == "Ours (Full)":
            fill = ours_fill
        elif row["Method Group"] == "External PPF":
            fill = external_fill
        else:
            fill = None
        if fill is not None:
            for col_idx in range(1, len(columns) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

    _auto_width(ws)

    meta = wb.create_sheet("Notes")
    meta["A1"] = "Stanford all-method comparison"
    meta["A2"] = "This workbook aggregates the root batch JSON files under experiments/results."
    meta["A3"] = "Internal Stanford methods and the 4 external PPF reproductions are listed together."
    meta["A4"] = "Same-backbone Baseline and Drost original PPF are both kept because they represent different reporting tracks."
    for cell_ref in ("A1", "A2", "A3", "A4"):
        meta[cell_ref].font = Font(name="Times New Roman", bold=(cell_ref == "A1"))
    meta.column_dimensions["A"].width = 120

    wb.save(XLSX_OUT)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    rows = [_aggregate_batch(spec) for spec in METHOD_SPECS]
    columns = [
        "Method Group",
        "Method",
        "n_cases",
        "Mean ADD",
        "Mean ADD-S",
        "SR@ADD<=0.02",
        "Mean Rotation Error (deg)",
        "Mean Translation Error",
        "Mean Registration Time (s)",
        "Mean Scene Preprocess Time (s)",
        "Mean Front-end Time (s)",
        "Mean Back-end Time (s)",
        "Candidate Inflation Mean",
        "Source Batch JSON",
    ]
    _write_csv(rows, columns)
    _write_workbook(rows, columns)
    print(f"Wrote {CSV_OUT}")
    print(f"Wrote {XLSX_OUT}")


if __name__ == "__main__":
    main()
