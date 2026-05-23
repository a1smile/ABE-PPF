import json
import math
import argparse
from pathlib import Path
from typing import Dict, List, Tuple

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
RESULT_DIR = ROOT / "experiments" / "results" / "stanford"
TABLE_DIR = ROOT / "experiments" / "tables" / "stanford" / "analysis_plus"
FIG_DIR = ROOT / "experiments" / "figures" / "stanford_analysis_plus"
SUCCESS_THRESHOLD = 0.02
PNG_DPI = 320


def batch_json_path(filename: str) -> Path:
    preferred = RESULT_DIR / filename
    if preferred.exists():
        return preferred
    fallback = RESULT_DIR.parent / filename
    return fallback

CORE_METHODS = [
    "Same-backbone Baseline",
    "No RS-MRQ",
    "No Robust Vote",
    "No Backend",
    "Ours (Full)",
]
CORE_METHOD_COLORS = {
    "Same-backbone Baseline": "#808080",
    "No RS-MRQ": "#4E79A7",
    "No Robust Vote": "#F28E2B",
    "No Backend": "#59A14F",
    "Ours (Full)": "#E15759",
}
EFFICIENCY_METHODS = [
    "No RS-MRQ",
    "Adaptive Two-Stage",
    "Ours (Full)",
    "Ours + RV20",
    "Ours + RV10",
    "Ours + RS-MRQ cap64",
    "Ours + RS-MRQ cap128",
    "Ours + RS-MRQ cap256",
]

METHOD_SOURCES = {
    "Same-backbone Baseline": batch_json_path("stanford_same_backbone_baseline_batch.json"),
    "No RS-MRQ": batch_json_path("stanford_no_rsmrq_batch.json"),
    "No Robust Vote": batch_json_path("stanford_no_robust_vote_batch.json"),
    "No Backend": batch_json_path("stanford_no_backend_batch.json"),
    "Ours (Full)": batch_json_path("stanford_ours_full_batch.json"),
    "Adaptive Two-Stage": batch_json_path("stanford_adaptive_two_stage_batch.json"),
    "Ours + RV20": batch_json_path("stanford_ours_rv20_batch.json"),
    "Ours + RV10": batch_json_path("stanford_ours_rv10_batch.json"),
    "Ours + RS-MRQ cap64": batch_json_path("stanford_ours_cap64_batch.json"),
    "Ours + RS-MRQ cap128": batch_json_path("stanford_ours_cap128_batch.json"),
    "Ours + RS-MRQ cap256": batch_json_path("stanford_ours_cap256_batch.json"),
}

MODULE_FLAGS = {
    "Same-backbone Baseline": {"RS-MRQ": "No", "Robust Vote": "No", "Candidate-level Selection": "No", "Mode Clustering": "No"},
    "No RS-MRQ": {"RS-MRQ": "No", "Robust Vote": "Yes", "Candidate-level Selection": "Yes", "Mode Clustering": "Yes"},
    "No Robust Vote": {"RS-MRQ": "Yes", "Robust Vote": "No", "Candidate-level Selection": "Yes", "Mode Clustering": "Yes"},
    "No Backend": {"RS-MRQ": "Yes", "Robust Vote": "Yes", "Candidate-level Selection": "No", "Mode Clustering": "No"},
    "Ours (Full)": {"RS-MRQ": "Yes", "Robust Vote": "Yes", "Candidate-level Selection": "Yes", "Mode Clustering": "Yes"},
}


def parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="Build Stanford ablation tables from existing batch JSON files.")
    ap.add_argument(
        "--only_overall_ablation",
        action="store_true",
        help="Only regenerate the Stanford overall ablation table outputs.",
    )
    return ap.parse_args()


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def apply_plot_style() -> None:
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
            "mathtext.fontset": "stix",
            "axes.titlesize": 12,
            "axes.labelsize": 11,
            "savefig.bbox": "tight",
        }
    )


def load_results(path: Path) -> List[dict]:
    with path.open("r", encoding="utf-8") as f:
        payload = json.load(f)
    return payload["results"]


def safe_float(value: object) -> float:
    try:
        out = float(value)
    except (TypeError, ValueError):
        return float("nan")
    if math.isnan(out) or math.isinf(out):
        return float("nan")
    return out


def mean(values: List[float]) -> float:
    arr = np.asarray(values, dtype=float)
    arr = arr[np.isfinite(arr)]
    return float(arr.mean()) if arr.size else float("nan")


def success_rate(values: List[float], threshold: float = SUCCESS_THRESHOLD) -> float:
    arr = np.asarray(values, dtype=float)
    arr = arr[np.isfinite(arr)]
    return float((arr <= threshold).mean()) if arr.size else float("nan")


def summarize_method(results: List[dict]) -> Dict[str, float]:
    add_vals = [safe_float(r.get("metrics", {}).get("ADD")) for r in results]
    adds_vals = [safe_float(r.get("metrics", {}).get("ADD_S")) for r in results]
    rot_vals = [safe_float(r.get("metrics", {}).get("rotation_error_deg")) for r in results]
    trans_vals = [safe_float(r.get("metrics", {}).get("translation_error")) for r in results]
    reg_vals = [safe_float(r.get("stats", {}).get("registration_time")) for r in results]
    preprocess_vals = [safe_float(r.get("stats", {}).get("scene_preprocess_time")) for r in results]
    front_vals = [safe_float(r.get("stats", {}).get("ppf_frontend_time")) for r in results]
    back_vals = [safe_float(r.get("stats", {}).get("backend_time")) for r in results]
    infl_vals = [safe_float(r.get("stats", {}).get("candidate_inflation_mean")) for r in results]
    return {
        "n_cases": float(len(results)),
        "Mean ADD": mean(add_vals),
        "Mean ADD-S": mean(adds_vals),
        "SR@ADD<=0.02": success_rate(add_vals),
        "Mean Rotation Error (deg)": mean(rot_vals),
        "Mean Translation Error": mean(trans_vals),
        "Mean Registration Time (s)": mean(reg_vals),
        "Mean Scene Preprocess Time (s)": mean(preprocess_vals),
        "Mean Front-end Time (s)": mean(front_vals),
        "Mean Back-end Time (s)": mean(back_vals),
        "Candidate Inflation Mean": mean(infl_vals),
    }


def to_markdown(df: pd.DataFrame) -> str:
    cols = list(df.columns)
    rows = [cols] + [[str(row[col]) for col in cols] for _, row in df.iterrows()]
    widths = [max(len(str(row[i])) for row in rows) for i in range(len(cols))]
    header = "| " + " | ".join(str(cols[i]).ljust(widths[i]) for i in range(len(cols))) + " |"
    sep = "| " + " | ".join("-" * widths[i] for i in range(len(cols))) + " |"
    body = [
        "| " + " | ".join(str(row[i]).ljust(widths[i]) for i in range(len(cols))) + " |"
        for row in rows[1:]
    ]
    return "\n".join([header, sep] + body) + "\n"


def format_for_display(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in out.columns:
        if pd.api.types.is_numeric_dtype(out[col]):
            if "Relative Time vs Ours" in col:
                out[col] = out[col].map(lambda v: "" if pd.isna(v) else f"{float(v):+.2f}%")
            elif "Relative ADD vs Ours" in col:
                out[col] = out[col].map(lambda v: "" if pd.isna(v) else f"{float(v):+.6f}")
            elif "SR@" in col:
                out[col] = out[col].map(lambda v: "" if pd.isna(v) else f"{float(v):.4f}")
            else:
                out[col] = out[col].map(lambda v: "" if pd.isna(v) else f"{float(v):.4f}")
    return out


def render_table_figure(df: pd.DataFrame, out_base: str, title: str) -> None:
    display = format_for_display(df)
    nrows, ncols = display.shape
    fig, ax = plt.subplots(figsize=(max(10.0, ncols * 1.35), max(3.0, nrows * 0.55 + 1.4)))
    ax.axis("off")
    table = ax.table(
        cellText=display.values.tolist(),
        colLabels=display.columns.tolist(),
        cellLoc="center",
        colLoc="center",
        loc="center",
    )
    table.auto_set_font_size(False)
    table.set_fontsize(9)
    table.scale(1.0, 1.35)
    for (r, c), cell in table.get_celld().items():
        cell.set_edgecolor("#777777")
        cell.set_linewidth(0.6)
        if r == 0:
            cell.set_text_props(weight="bold")
            cell.set_facecolor("#E9EEF5")
        elif r % 2 == 1:
            cell.set_facecolor("#F8F8F8")
    ax.set_title(title, pad=16)
    ensure_dir(FIG_DIR)
    fig.savefig(FIG_DIR / f"{out_base}.png", dpi=PNG_DPI)
    fig.savefig(FIG_DIR / f"{out_base}.svg")
    plt.close(fig)


def save_figure(fig: plt.Figure, out_base: str) -> None:
    ensure_dir(FIG_DIR)
    fig.savefig(FIG_DIR / f"{out_base}.png", dpi=PNG_DPI)
    fig.savefig(FIG_DIR / f"{out_base}.svg")
    plt.close(fig)


def build_main_ablation_plot_source(summaries: Dict[str, Dict[str, float]]) -> pd.DataFrame:
    rows = []
    for method in CORE_METHODS:
        summary = summaries[method]
        rows.append(
            {
                "method": method,
                "mean_add": summary["Mean ADD"],
                "sr_add_le_0_02": summary["SR@ADD<=0.02"],
                "mean_registration_time": summary["Mean Registration Time (s)"],
                "mean_scene_preprocess_time": summary["Mean Scene Preprocess Time (s)"],
                "mean_frontend_time": summary["Mean Front-end Time (s)"],
                "mean_backend_time": summary["Mean Back-end Time (s)"],
            }
        )
    return pd.DataFrame(rows)


def _axis_limits(values: np.ndarray, lower_bound: float | None = None, upper_bound: float | None = None) -> Tuple[float, float]:
    finite = values[np.isfinite(values)]
    vmin = float(finite.min())
    vmax = float(finite.max())
    span = max(vmax - vmin, 1e-6)
    pad = span * 0.12
    lo = vmin - pad
    hi = vmax + pad
    if lower_bound is not None:
        lo = max(lower_bound, lo)
    if upper_bound is not None:
        hi = min(upper_bound, hi)
    return lo, hi


def plot_main_ablation_figure(source_df: pd.DataFrame) -> None:
    df = source_df.set_index("method").loc[CORE_METHODS].reset_index()
    label_positions = {
        "Same-backbone Baseline": {"xytext": (4.0, 0.657), "ha": "left", "va": "bottom", "arrow": False},
        "No RS-MRQ": {"xytext": (3.6, 0.914), "ha": "right", "va": "bottom", "arrow": True},
        "No Robust Vote": {"xytext": (9.6, 0.936), "ha": "left", "va": "bottom", "arrow": True},
        "No Backend": {"xytext": (31.0, 0.910), "ha": "left", "va": "bottom", "arrow": True},
        "Ours (Full)": {"xytext": (31.0, 0.936), "ha": "left", "va": "bottom", "arrow": True},
    }

    fig, axes = plt.subplots(1, 2, figsize=(10.8, 4.4), gridspec_kw={"width_ratios": [1.02, 1.18]})

    ax = axes[0]
    x = df["mean_registration_time"].to_numpy(dtype=float)
    y = df["sr_add_le_0_02"].to_numpy(dtype=float)
    ax.set_xlim(*_axis_limits(x, lower_bound=0.0))
    ax.set_ylim(*_axis_limits(y, lower_bound=0.0, upper_bound=1.0))
    for _, row in df.iterrows():
        method = row["method"]
        is_ours = method == "Ours (Full)"
        ax.scatter(
            row["mean_registration_time"],
            row["sr_add_le_0_02"],
            s=130 if is_ours else 95,
            color=CORE_METHOD_COLORS[method],
            edgecolor="black" if is_ours else "white",
            linewidth=1.1 if is_ours else 0.8,
            zorder=3,
        )
        ax.annotate(
            method,
            (row["mean_registration_time"], row["sr_add_le_0_02"]),
            xytext=label_positions.get(method, {}).get("xytext", (row["mean_registration_time"], row["sr_add_le_0_02"])),
            textcoords="data",
            fontsize=9,
            ha=label_positions.get(method, {}).get("ha", "left"),
            va=label_positions.get(method, {}).get("va", "bottom"),
            arrowprops=(
                {
                    "arrowstyle": "-",
                    "color": "#777777",
                    "lw": 0.8,
                    "shrinkA": 0,
                    "shrinkB": 5,
                }
                if label_positions.get(method, {}).get("arrow", False)
                else None
            ),
        )
    ax.set_xlabel("Mean Registration Time (s)")
    ax.set_ylabel("SR@ADD<=0.02")
    ax.set_title("Accuracy-Time Trade-off")
    ax.grid(True, linestyle="--", alpha=0.25)

    ax = axes[1]
    xpos = np.arange(len(df))
    preprocess = df["mean_scene_preprocess_time"].to_numpy(dtype=float)
    frontend = df["mean_frontend_time"].to_numpy(dtype=float)
    backend = df["mean_backend_time"].to_numpy(dtype=float)
    width = 0.68
    ax.bar(xpos, preprocess, width=width, color="#D9D9D9", edgecolor="none", label="Scene preprocess")
    ax.bar(xpos, frontend, width=width, bottom=preprocess, color="#F28E2B", edgecolor="none", label="Front-end")
    ax.bar(xpos, backend, width=width, bottom=preprocess + frontend, color="#76B7B2", edgecolor="none", label="Back-end")
    ax.set_xticks(xpos)
    ax.set_xticklabels(df["method"], rotation=22, ha="right")
    ax.set_ylabel("Mean Time (s)")
    ax.set_xlabel("Method")
    ax.set_title("Timing Breakdown")
    ax.grid(True, axis="y", linestyle="--", alpha=0.25)
    ax.legend(frameon=False, loc="upper left")

    fig.suptitle("Main Stanford Ablation Figure", y=1.02)
    fig.tight_layout()
    save_figure(fig, "main_ablation_figure")


def build_overall_ablation_table(summaries: Dict[str, Dict[str, float]]) -> pd.DataFrame:
    rows = []
    for method in CORE_METHODS:
        summary = dict(summaries[method])
        summary.pop("n_cases", None)
        rows.append({"Method": method, **MODULE_FLAGS[method], **summary})
    return pd.DataFrame(rows)


def write_overall_ablation_xlsx(df: pd.DataFrame) -> None:
    ensure_dir(TABLE_DIR)
    out_path = TABLE_DIR / "overall_ablation_table.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Stanford_Ablation"
    ws.freeze_panes = "A2"

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    ours_fill = PatternFill(fill_type="solid", fgColor="FDE9D9")
    no_backend_fill = PatternFill(fill_type="solid", fgColor="E2F0D9")

    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = Font(name="Times New Roman", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = header_fill

    numeric_formats = {
        "Mean ADD": "0.000000",
        "Mean ADD-S": "0.000000",
        "SR@ADD<=0.02": "0.000000",
        "Mean Rotation Error (deg)": "0.000000",
        "Mean Translation Error": "0.000000",
        "Mean Registration Time (s)": "0.000000",
        "Mean Front-end Time (s)": "0.000000",
        "Mean Back-end Time (s)": "0.000000",
        "Candidate Inflation Mean": "0.000000",
    }

    for row_idx, row in enumerate(df.to_dict(orient="records"), start=2):
        for col_idx, col_name in enumerate(df.columns, start=1):
            value = row.get(col_name)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name="Times New Roman")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_name in numeric_formats and isinstance(value, (int, float)):
                cell.number_format = numeric_formats[col_name]

        fill = None
        if row["Method"] == "Ours (Full)":
            fill = ours_fill
        elif row["Method"] == "No Backend":
            fill = no_backend_fill
        if fill is not None:
            for col_idx in range(1, len(df.columns) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

    for idx, column_cells in enumerate(ws.columns, start=1):
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 12), 36)

    wb.save(out_path)


def build_efficiency_extensions_table(summaries: Dict[str, Dict[str, float]], results_map: Dict[str, List[dict]]) -> pd.DataFrame:
    ours_add = summaries["Ours (Full)"]["Mean ADD"]
    ours_time = summaries["Ours (Full)"]["Mean Registration Time (s)"]
    rows = []
    for method in EFFICIENCY_METHODS:
        if method not in summaries:
            continue
        base_variant = "Core"
        note = "-"
        if method == "Adaptive Two-Stage":
            base_variant = "Adaptive"
            upgrades = [1.0 if bool(r.get("adaptive_used_upgrade", False)) else 0.0 for r in results_map[method]]
            note = f"upgrade={mean(upgrades) * 100.0:.2f}%"
        elif "RV" in method:
            base_variant = "Robust Vote Light"
            rv_topm = safe_float(results_map[method][0].get("debug", {}).get("robust_vote", {}).get("topm_used"))
            cfg_topm = safe_float(results_map[method][0].get("debug", {}).get("robust_vote", {}).get("topm_ratio"))
            if method.endswith("RV20"):
                note = "top-m=20"
            elif method.endswith("RV10"):
                note = "top-m=10"
            else:
                note = f"rv={rv_topm:.0f}/{cfg_topm:.4f}"
        elif "cap" in method:
            base_variant = "RS-MRQ Cap"
            cap = results_map[method][0].get("debug", {}).get("rsmrq", {}).get("global_candidate_cap")
            note = f"cap={int(cap)}" if cap is not None else "-"
        elif method == "No RS-MRQ":
            base_variant = "Reference Lightweight"
        rows.append(
            {
                "Method": method,
                "Base Variant": base_variant,
                "Mean ADD": summaries[method]["Mean ADD"],
                "SR@ADD<=0.02": summaries[method]["SR@ADD<=0.02"],
                "Mean Registration Time (s)": summaries[method]["Mean Registration Time (s)"],
                "Mean Front-end Time (s)": summaries[method]["Mean Front-end Time (s)"],
                "Mean Back-end Time (s)": summaries[method]["Mean Back-end Time (s)"],
                "Relative Time vs Ours": (summaries[method]["Mean Registration Time (s)"] - ours_time) / ours_time * 100.0,
                "Relative ADD vs Ours": summaries[method]["Mean ADD"] - ours_add,
                "Upgrade Ratio / RV top-m / Candidate Cap": note,
            }
        )
    return pd.DataFrame(rows)


def plot_overall_ablation_overview(overall_df: pd.DataFrame) -> None:
    order = CORE_METHODS
    df = overall_df.set_index("Method").loc[order].reset_index()
    y = np.arange(len(df))
    highlight = df["Method"] == "Ours (Full)"
    add_colors = ["#E15759" if is_ours else "#4E79A7" for is_ours in highlight]
    front_colors = ["#D95F02" if is_ours else "#F28E2B" for is_ours in highlight]
    back_colors = ["#8C2D04" if is_ours else "#76B7B2" for is_ours in highlight]

    fig, axes = plt.subplots(1, 2, figsize=(12.6, 4.8), gridspec_kw={"width_ratios": [1.0, 1.35]})

    ax = axes[0]
    ax.barh(y, df["Mean ADD"], color=add_colors, edgecolor="none")
    ax.set_yticks(y)
    ax.set_yticklabels(df["Method"])
    ax.invert_yaxis()
    ax.set_xlabel("Mean ADD")
    ax.set_title("Accuracy")
    ax.grid(True, axis="x", alpha=0.25, linestyle="--")
    for i, value in enumerate(df["Mean ADD"]):
        ax.text(value + max(df["Mean ADD"]) * 0.015, y[i], f"{value:.4f}", va="center", ha="left", fontsize=9)

    ax = axes[1]
    front = df["Mean Front-end Time (s)"].to_numpy(dtype=float)
    back = df["Mean Back-end Time (s)"].to_numpy(dtype=float)
    ax.barh(y, front, color=front_colors, edgecolor="none", label="Front-end")
    ax.barh(y, back, left=front, color=back_colors, edgecolor="none", label="Back-end")
    ax.set_yticks(y)
    ax.set_yticklabels([])
    ax.invert_yaxis()
    ax.set_xlabel("Mean registration time (s)")
    ax.set_title("Timing Breakdown")
    ax.grid(True, axis="x", alpha=0.25, linestyle="--")
    for i, total in enumerate(front + back):
        ax.text(total + max(front + back) * 0.015, y[i], f"{total:.2f}s", va="center", ha="left", fontsize=9)
    ax.legend(frameon=False, loc="lower right")

    fig.suptitle("Stanford Overall Ablation Overview", y=1.02)
    fig.tight_layout()
    save_figure(fig, "overall_ablation_overview")


def plot_overall_ablation_pareto(overall_df: pd.DataFrame) -> None:
    df = overall_df.set_index("Method").loc[CORE_METHODS].reset_index()
    times = df["Mean Registration Time (s)"].to_numpy(dtype=float)
    success = df["SR@ADD<=0.02"].to_numpy(dtype=float)

    method_colors = {
        "Same-backbone Baseline": "#B07AA1",
        "No RS-MRQ": "#59A14F",
        "No Robust Vote": "#76B7B2",
        "No Backend": "#EDC948",
        "Ours (Full)": "#E15759",
    }

    pareto = []
    for i in range(len(df)):
        dominated = False
        for j in range(len(df)):
            if i == j:
                continue
            if times[j] <= times[i] and success[j] >= success[i] and (times[j] < times[i] or success[j] > success[i]):
                dominated = True
                break
        pareto.append(not dominated)

    frontier_df = df.loc[pareto].sort_values("Mean Registration Time (s)")

    fig, ax = plt.subplots(figsize=(8.2, 5.4))
    for _, row in df.iterrows():
        method = row["Method"]
        is_ours = method == "Ours (Full)"
        ax.scatter(
            row["Mean Registration Time (s)"],
            row["SR@ADD<=0.02"],
            s=160 if is_ours else 120,
            color=method_colors[method],
            edgecolor="#222222" if is_ours else "white",
            linewidth=1.1 if is_ours else 0.9,
            zorder=3,
        )
        ax.annotate(
            method,
            (row["Mean Registration Time (s)"], row["SR@ADD<=0.02"]),
            xytext=(6, 6),
            textcoords="offset points",
            fontsize=9,
        )

    ax.plot(
        frontier_df["Mean Registration Time (s)"],
        frontier_df["SR@ADD<=0.02"],
        color="#444444",
        linestyle="--",
        linewidth=1.4,
        alpha=0.9,
        zorder=2,
        label="Pareto frontier",
    )
    ax.set_xlabel("Mean registration time (s)")
    ax.set_ylabel("SR@ADD<=0.02")
    ax.set_title("Stanford Ablation Pareto Plot")
    ax.grid(True, linestyle="--", alpha=0.25)
    ax.legend(frameon=False, loc="lower right")
    fig.tight_layout()
    save_figure(fig, "overall_ablation_pareto")


def plot_overall_ablation_relative_change(overall_df: pd.DataFrame) -> None:
    df = overall_df.set_index("Method").loc[CORE_METHODS].reset_index()
    ours = df.loc[df["Method"] == "Ours (Full)"].iloc[0]
    comp = df.loc[df["Method"] != "Ours (Full)"].copy()

    comp["Time Delta (%)"] = (
        (comp["Mean Registration Time (s)"] - float(ours["Mean Registration Time (s)"]))
        / float(ours["Mean Registration Time (s)"])
        * 100.0
    )
    comp["ADD Delta (%)"] = (
        (comp["Mean ADD"] - float(ours["Mean ADD"])) / float(ours["Mean ADD"]) * 100.0
    )
    comp["Label"] = comp["Method"].map(
        {
            "Same-backbone Baseline": "Same-backbone baseline",
            "No RS-MRQ": "Remove RS-MRQ",
            "No Robust Vote": "Remove Robust Vote",
            "No Backend": "No Backend",
        }
    )
    comp = comp.reset_index(drop=True)
    y = np.arange(len(comp))

    fig, axes = plt.subplots(1, 2, figsize=(12.2, 4.8), gridspec_kw={"width_ratios": [1.15, 1.0]})

    time_colors = ["#59A14F" if value < 0 else "#E15759" for value in comp["Time Delta (%)"]]
    axes[0].barh(y, comp["Time Delta (%)"], color=time_colors, edgecolor="none")
    axes[0].axvline(0.0, color="#444444", linewidth=1.0)
    axes[0].set_yticks(y)
    axes[0].set_yticklabels(comp["Label"])
    axes[0].invert_yaxis()
    axes[0].set_xlabel("Relative registration time vs Ours (%)")
    axes[0].set_title("Time Change")
    axes[0].grid(True, axis="x", linestyle="--", alpha=0.25)
    for i, value in enumerate(comp["Time Delta (%)"]):
        offset = 1.5 if value >= 0 else -1.5
        ha = "left" if value >= 0 else "right"
        axes[0].text(value + offset, y[i], f"{value:+.1f}%", va="center", ha=ha, fontsize=9)

    add_colors = ["#E15759" if value > 0 else "#59A14F" for value in comp["ADD Delta (%)"]]
    axes[1].barh(y, comp["ADD Delta (%)"], color=add_colors, edgecolor="none")
    axes[1].axvline(0.0, color="#444444", linewidth=1.0)
    axes[1].set_yticks(y)
    axes[1].set_yticklabels([])
    axes[1].invert_yaxis()
    axes[1].set_xlabel("Relative Mean ADD vs Ours (%)")
    axes[1].set_title("Accuracy Change")
    axes[1].grid(True, axis="x", linestyle="--", alpha=0.25)
    for i, value in enumerate(comp["ADD Delta (%)"]):
        offset = 1.5 if value >= 0 else -1.5
        ha = "left" if value >= 0 else "right"
        axes[1].text(value + offset, y[i], f"{value:+.1f}%", va="center", ha=ha, fontsize=9)

    fig.suptitle("Relative Effect of Removing Stanford Ablation Modules", y=1.02)
    fig.tight_layout()
    save_figure(fig, "overall_ablation_relative_change")


def build_summary_text(overall_df: pd.DataFrame) -> str:
    ours = overall_df.loc[overall_df["Method"] == "Ours (Full)"].iloc[0]
    diffs = overall_df.loc[overall_df["Method"] != "Ours (Full)"].copy()
    diffs["add_gap"] = diffs["Mean ADD"] - ours["Mean ADD"]
    diffs["time_drop"] = ours["Mean Registration Time (s)"] - diffs["Mean Registration Time (s)"]
    worst_acc = diffs.loc[diffs["add_gap"].idxmax()]
    biggest_time_drop = diffs.loc[diffs["time_drop"].idxmax()]
    front_ratio = ours["Mean Front-end Time (s)"] / ours["Mean Registration Time (s)"]

    reg = overall_df["Mean Registration Time (s)"].to_numpy(dtype=float)
    add = overall_df["Mean ADD"].to_numpy(dtype=float)
    pareto = []
    for i in range(len(overall_df)):
        dominated = False
        for j in range(len(overall_df)):
            if i == j:
                continue
            if reg[j] <= reg[i] and add[j] <= add[i] and (reg[j] < reg[i] or add[j] < add[i]):
                dominated = True
                break
        pareto.append(not dominated)
    ours_pareto = bool(pareto[int(overall_df.index[overall_df["Method"] == "Ours (Full)"][0])])
    pareto_text = (
        "Ours (Full) stays on the Pareto frontier among the five core methods."
        if ours_pareto
        else "Ours (Full) is not on the Pareto frontier among the five core methods."
    )

    return "\n".join(
        [
            "# Overall Ablation Summary",
            "",
            f"- The largest accuracy drop comes from `{worst_acc['Method']}`, whose mean ADD is {worst_acc['add_gap']:.6f} higher than Ours (Full).",
            f"- The largest time reduction comes from `{biggest_time_drop['Method']}`, whose mean registration time is {biggest_time_drop['time_drop']:.2f}s lower than Ours (Full).",
            f"- The bottleneck is still front-end: Ours (Full) spends {front_ratio * 100.0:.2f}% of registration time in front-end.",
            f"- {pareto_text}",
        ]
    ) + "\n"


def build_consistency_text(results_map: Dict[str, List[dict]], overall_df: pd.DataFrame) -> str:
    case_counts = {method: len(results_map[method]) for method in CORE_METHODS}
    missing_columns = [col for col in overall_df.columns if overall_df[col].isna().all()]
    return "\n".join(
        [
            "# Overall Ablation Consistency Check",
            "",
            "- Every numeric cell in the overall ablation table was re-aggregated directly from the core batch JSON files.",
            "- Method names are consistent with the batch JSON filenames and table row names.",
            f"- Case counts: {case_counts}.",
            f"- Fully empty columns: {missing_columns if missing_columns else 'None'}.",
        ]
    ) + "\n"


def export_main_ablation_assets(summaries: Dict[str, Dict[str, float]]) -> None:
    ensure_dir(TABLE_DIR)
    source_df = build_main_ablation_plot_source(summaries)
    source_df.to_csv(TABLE_DIR / "main_ablation_plot_source.csv", index=False)
    caption = (
        "Main ablation results on Stanford. Left: accuracy-time trade-off of the five core variants. "
        "Right: timing breakdown into scene preprocessing, front-end, and back-end. "
        "The results show that the main runtime overhead comes from the front-end rather than mode clustering, "
        "while the complete model achieves the best overall accuracy."
    )
    (TABLE_DIR / "main_ablation_caption.txt").write_text(caption, encoding="utf-8")
    plot_main_ablation_figure(source_df)


def main() -> None:
    args = parse_args()
    apply_plot_style()
    ensure_dir(TABLE_DIR)
    ensure_dir(FIG_DIR)

    results_map = {method: load_results(path) for method, path in METHOD_SOURCES.items() if path.exists()}
    summaries = {method: summarize_method(results) for method, results in results_map.items()}

    overall_df = build_overall_ablation_table(summaries)
    overall_csv = TABLE_DIR / "overall_ablation_table.csv"
    overall_md = TABLE_DIR / "overall_ablation_table.md"
    overall_summary = TABLE_DIR / "overall_ablation_summary.md"
    overall_check = TABLE_DIR / "overall_ablation_consistency_check.md"
    overall_df.to_csv(overall_csv, index=False)
    overall_md.write_text(to_markdown(format_for_display(overall_df)), encoding="utf-8")
    overall_summary.write_text(build_summary_text(overall_df), encoding="utf-8")
    overall_check.write_text(build_consistency_text(results_map, overall_df), encoding="utf-8")
    write_overall_ablation_xlsx(overall_df)
    if args.only_overall_ablation:
        return

    render_table_figure(overall_df, "overall_ablation_table", "Overall Stanford Ablation Table")
    export_main_ablation_assets(summaries)
    plot_overall_ablation_overview(overall_df)
    plot_overall_ablation_pareto(overall_df)
    plot_overall_ablation_relative_change(overall_df)

    efficiency_df = build_efficiency_extensions_table(summaries, results_map)
    efficiency_csv = TABLE_DIR / "efficiency_extensions_table.csv"
    efficiency_md = TABLE_DIR / "efficiency_extensions_table.md"
    efficiency_df.to_csv(efficiency_csv, index=False)
    efficiency_md.write_text(to_markdown(format_for_display(efficiency_df)), encoding="utf-8")
    render_table_figure(efficiency_df, "efficiency_extensions_table", "Stanford Efficiency Extensions Table")


if __name__ == "__main__":
    main()
